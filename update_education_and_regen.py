import json, re
from pathlib import Path
from openpyxl import load_workbook
from exporter import export_payload
from cv_report import render_personnel_docx

OUT_DIR = Path(r"C:\hscan\cv_outputs")
payload_path = OUT_DIR / "payload_daftar_tenaga_ahli.json"
backup_path = OUT_DIR / "payload_daftar_tenaga_ahli_before_pendidikan_short.json"
xlsx_path = OUT_DIR / "Daftar_Tenaga_Ahli_Jawa_1_pendidikan_singkat.xlsx"
docx_path = OUT_DIR / "Daftar_Tenaga_Ahli_Jawa_1_pendidikan_singkat.docx"

payload = json.loads(payload_path.read_text(encoding="utf-8"))
if not backup_path.exists():
    backup_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

DEGREE_PATTERNS = [
    (r"\bS\s*[-.]?\s*1\b", "S1"),
    (r"\bS\s*[-.]?\s*2\b", "S2"),
    (r"\bS\s*[-.]?\s*3\b", "S3"),
    (r"\bD\s*[-.]?\s*4\b", "D4"),
    (r"\bD\s*[-.]?\s*3\b", "D3"),
    (r"\bSMK\b|\bSekolah Menengah Kejuruan\b", "SMK"),
    (r"\bSMA\b|\bSekolah Menengah Atas\b", "SMA"),
]
STOP_WORDS = [
    "Universitas", "Institut", "Sekolah Tinggi", "Politeknik", "Yayasan", "STT", "STIE",
    "Negeri", "Swasta", "Muhammadiyah", "Islam", "Atma", "Jaya", "Bina", "Darma",
]

# Normalize common OCR junk and remove institution/year; keep only degree + discipline/major.
def shorten_education(text: str) -> str:
    original = str(text or "")
    s = re.sub(r"^[^A-Za-z0-9]*(?:i\s+)?", "", original).strip()
    s = s.replace('Teknik Konstruksi Gedung', 'Teknik Sipil')
    s = re.sub(r"\s+", " ", s)
    degree = ""
    start = 0
    for pat, repl in DEGREE_PATTERNS:
        m = re.search(pat, s, flags=re.I)
        if m:
            degree = repl
            start = m.end()
            break
    if degree:
        rest = s[start:].strip(" -.,:")
    else:
        # vocational lines may start with major only
        rest = s
    # Cut at institution indicators or year.
    cut_positions = []
    for w in STOP_WORDS:
        m = re.search(r"\b" + re.escape(w) + r"\b", rest, flags=re.I)
        if m:
            cut_positions.append(m.start())
    m = re.search(r"\b(19|20)\d{2}\b", rest)
    if m:
        cut_positions.append(m.start())
    if cut_positions:
        rest = rest[:min(cut_positions)]
    rest = rest.strip(" -.,:")
    # If rest has too many trailing institution/location fragments, keep up to known major phrase.
    known = [
        "Teknik Sipil", "Teknik Elektro", "Teknik Arsitektur", "Arsitektur", "Teknik Kimia",
        "Teknik Mesin", "Teknik Kendaraan Ringan", "Teknik Pemanfaatan Tenaga Listrik",
        "Teknik Sepeda Motor", "Akuntansi", "Administrasi", "Konstruksi Gedung",
    ]
    for k in known:
        if re.search(re.escape(k), rest, flags=re.I):
            rest = k
            break
    if degree and rest:
        return f"{degree} {rest}".strip()
    if degree:
        return degree
    return rest or original

for person in payload["personel"]:
    person["kualifikasi_pendidikan"] = shorten_education(person.get("kualifikasi_pendidikan", ""))

payload_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

# Excel via cloned script exporter.
export_result = export_payload("daftar_tenaga_ahli", payload, xlsx_path)

# Word generation mirroring previous layout.
render_personnel_docx(payload, docx_path)

wb = load_workbook(xlsx_path, read_only=True, data_only=True)
ws = wb.active
samples = [ws.cell(r,5).value for r in range(6, min(ws.max_row, 15)+1)]
print(json.dumps({
    "success": True,
    "xlsx": str(xlsx_path),
    "docx": str(docx_path),
    "payload": str(payload_path),
    "rows": len(payload['personel']),
    "excel_rows": ws.max_row - 5,
    "sample_pendidikan": samples,
    "export_result": export_result,
}, ensure_ascii=False, indent=2))
wb.close()
