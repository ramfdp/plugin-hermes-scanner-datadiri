import argparse
import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


def text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


def _add_table(document, rows):
    table = document.add_table(rows=0, cols=len(rows[0]))
    table.style = "Table Grid"
    for row in rows:
        for cell, value in zip(table.add_row().cells, row):
            cell.text = text(value)
    return table


def render_report(payload: dict, output_path: Path) -> dict:
    required = ("cv_file", "biodata", "experience_validation", "attachment_cross_check", "findings")
    if not isinstance(payload, dict):
        raise ValueError("Payload harus berupa object.")
    missing = [key for key in required if key not in payload]
    if missing:
        raise ValueError("Field wajib belum tersedia: " + ", ".join(missing))
<<<<<<< HEAD
    for key in ("experience_validation", "attachment_cross_check", "employer_validation"):
        rows = payload.get(key, [])
        if not isinstance(rows, list) or any(not isinstance(row, dict) for row in rows):
            raise ValueError(f"'{key}' harus berupa array object.")
    for key in ("findings", "internet_sources", "chronology_validation"):
        values = payload.get(key, [])
        if not isinstance(values, list) or any(not isinstance(value, str) for value in values):
            raise ValueError(f"'{key}' harus berupa array string.")
    if not isinstance(payload.get("source_artifacts", {}), dict):
        raise ValueError("source_artifacts harus berupa object.")
    for employer in payload.get("employer_validation", []):
        sources = employer.get("sources", [])
        if not isinstance(sources, list) or any(not isinstance(url, str) for url in sources):
            raise ValueError("sources perusahaan harus berupa array URL.")
        if employer.get("status") in ("terverifikasi", "sebagian terverifikasi") and not any(
            url.startswith(("https://", "http://")) for url in sources
        ):
            raise ValueError("Status perusahaan terverifikasi memerlukan URL sumber.")

    parts = [paragraph("LAPORAN REVIEW CV", "Title"), paragraph("Dokumen hasil pemetaan biodata, validasi pengalaman, dan cross-check lampiran."), paragraph(f"File CV: {payload['cv_file']}")]
    if payload.get("kak_file"):
        parts.append(paragraph(f"File KAK: {payload['kak_file']}"))

    parts.append(paragraph("1. Pemetaan Biodata", "Heading1"))
    biodata = payload["biodata"] if isinstance(payload["biodata"], dict) else {"biodata": payload["biodata"]}
    parts.append(key_value_table(biodata))

    parts.append(paragraph("2. Kesesuaian Pengalaman dengan Posisi dalam KAK", "Heading1"))
=======
    for key in ("experience_validation", "attachment_cross_check", "findings", "internet_sources"):
        value = payload.get(key)
        if value is not None and not isinstance(value, list):
            raise ValueError(f"'{key}' harus berupa array/list.")
>>>>>>> 47b178571b83a89922cbff90af5c0b5f00dff3c2
    experience = payload["experience_validation"] or []
    cross_checks = payload["attachment_cross_check"] or []
    if any(not isinstance(item, dict) for item in [*experience, *cross_checks]):
        raise ValueError("Item validasi pengalaman dan cross-check harus berupa object.")

    document = Document()
    document.add_paragraph("LAPORAN REVIEW CV", "Title")
    document.add_paragraph("Dokumen hasil pemetaan biodata, validasi pengalaman, dan cross-check lampiran.")
    document.add_paragraph(f"File CV: {payload['cv_file']}")
    if payload.get("kak_file"):
        document.add_paragraph(f"File KAK: {payload['kak_file']}")

    document.add_heading("1. Pemetaan Biodata", 1)
    biodata = payload["biodata"] if isinstance(payload["biodata"], dict) else {"biodata": payload["biodata"]}
    _add_table(document, [["Field", "Nilai"], *[
        [key.replace("_", " ").title(), value] for key, value in biodata.items()
    ]])

    document.add_heading("2. Kesesuaian Pengalaman dengan Posisi dalam KAK", 1)
    rows = [["Posisi/Kriteria KAK", "Klaim CV", "Validasi Internet", "Status", "Catatan"]]
    rows.extend([[item.get("kak_requirement", ""), item.get("cv_claim", ""),
                  item.get("internet_validation", ""), item.get("status", "belum diverifikasi"),
                  item.get("notes", "")] for item in experience])
    _add_table(document, rows if experience else [["Tidak ada pengalaman yang dapat divalidasi."]])

    document.add_heading("3. Cross-check CV dan Lampiran", 1)
    rows = [["Lampiran", "Field", "CV", "Lampiran", "Status/Catatan"]]
    rows.extend([[item.get("attachment", ""), item.get("field", ""), item.get("cv_value", ""),
                  item.get("attachment_value", ""), item.get("status", "")] for item in cross_checks])
    _add_table(document, rows if cross_checks else [["Tidak ada lampiran yang dicross-check."]])

<<<<<<< HEAD
    parts.append(paragraph("4. Temuan/Janggal yang Perlu Dikomentari", "Heading1"))
    parts.extend(paragraph(value, "ListBullet") for value in (payload["findings"] or ["Tidak ada temuan."]))
    parts.append(paragraph("5. Sumber Validasi Internet", "Heading1"))
    parts.extend(paragraph(value, "ListBullet") for value in (payload.get("internet_sources") or ["Tidak ada sumber yang dicatat."]))
    parts.append(paragraph("6. Kesimpulan", "Heading1"))
    parts.append(paragraph(payload.get("conclusion", "Belum ada kesimpulan.")))
    parts.append(paragraph("Status akhir: " + text(payload.get("status", "perlu review manual"))))

    if "employer_validation" in payload:
        parts.append(paragraph("7. Verifikasi Perusahaan", "Heading1"))
        rows = [["Perusahaan", "Status", "Bukti dan Batasan", "Sumber", "Tanggal Pemeriksaan"]]
        rows.extend([[item.get("employer", ""), item.get("status", "belum dapat diverifikasi"),
                      item.get("evidence", ""), "\n".join(item.get("sources", [])),
                      item.get("checked_at", "")] for item in payload["employer_validation"]])
        parts.append(table(rows))
        parts.append(paragraph("Keberadaan perusahaan bukan bukti bahwa kandidat pernah bekerja di sana."))
    if "chronology_validation" in payload:
        parts.append(paragraph("8. Validasi Kronologi", "Heading1"))
        parts.extend(paragraph(item) for item in payload["chronology_validation"])
    if payload.get("source_artifacts"):
        parts.append(paragraph("9. Dokumen Sumber", "Heading1"))
        parts.append(key_value_table(payload["source_artifacts"]))

    document_xml = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>' + "".join(parts) + '<w:sectPr/></w:body></w:document>'
    styles_xml = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:styles xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:style w:type="paragraph" w:styleId="Normal"><w:name w:val="Normal"/></w:style><w:style w:type="paragraph" w:styleId="Title"><w:name w:val="Title"/></w:style><w:style w:type="paragraph" w:styleId="Heading1"><w:name w:val="heading 1"/></w:style><w:style w:type="paragraph" w:styleId="ListBullet"><w:name w:val="List Bullet"/></w:style></w:styles>'
    content_types = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/><Override PartName="/word/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/></Types>'
    rels = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>'
    document_rels = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/></Relationships>'

    output_path.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(output_path, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", content_types)
        archive.writestr("_rels/.rels", rels)
        archive.writestr("word/document.xml", document_xml)
        archive.writestr("word/styles.xml", styles_xml)
        archive.writestr("word/_rels/document.xml.rels", document_rels)
    json_path = output_path.with_suffix(".json")
    json_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return {"success": True, "template": "cv_review", "output_file": str(output_path.resolve()),
            "json_file": str(json_path.resolve()), "finding_count": len(payload["findings"] or [])}
=======
    document.add_heading("4. Temuan/Janggal yang Perlu Dikomentari", 1)
    for value in payload["findings"] or ["Tidak ada temuan."]:
        document.add_paragraph(text(value), "List Bullet")
    document.add_heading("5. Sumber Validasi Internet", 1)
    for value in payload.get("internet_sources") or ["Tidak ada sumber yang dicatat."]:
        document.add_paragraph(text(value), "List Bullet")
    document.add_heading("6. Kesimpulan", 1)
    document.add_paragraph(text(payload.get("conclusion", "Belum ada kesimpulan.")))
    document.add_paragraph("Status akhir: " + text(payload.get("status", "perlu review manual")))

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return {"success": True, "template": "cv_review", "output_file": str(output_path.resolve()),
            "finding_count": len(payload["findings"] or [])}
>>>>>>> 47b178571b83a89922cbff90af5c0b5f00dff3c2


def render_personnel_docx(payload: dict, output_path: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = section.bottom_margin = Inches(0.35)
    section.left_margin = section.right_margin = Inches(0.25)
    titles = [(payload.get("judul", "DAFTAR TENAGA AHLI"), 11),
              (payload.get("wilayah", ""), 10), (payload.get("pekerjaan", ""), 10)]
    for value, size in titles:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(text(value))
        run.bold, run.font.name, run.font.size = True, "Century Gothic", Pt(size)

    headers = ["No", "Nama Personel", "NIK", "Jabatan Personel", "Kualifikasi Pendidikan",
               "Sertifikat Keahlian", "Pengalaman min dalam KAK (Tahun)",
               "Pengalaman Kerja (Bulan)", "Pengalaman Kerja (Tahun)"]
    table = document.add_table(rows=1, cols=len(headers))
    table.style, table.autofit = "Table Grid", True

    def set_cell(cell, value, *, bold=False, size=7, align=WD_ALIGN_PARAGRAPH.CENTER):
        cell.text = ""
        paragraph = cell.paragraphs[0]
        paragraph.alignment = align
        run = paragraph.add_run(text(value))
        run.bold, run.font.name, run.font.size = bold, "Century Gothic", Pt(size)

    for cell, header in zip(table.rows[0].cells, headers):
        set_cell(cell, header, bold=True)
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "8EA9D8")
        cell._tc.get_or_add_tcPr().append(shading)
    for number, person in enumerate(payload["personel"], 1):
        years = person.get("pengalaman_kerja_tahun", "")
        years_display = "" if years in (None, "") else f"{float(years):.2f}".replace(".", ",")
        min_kak = person.get("pengalaman_min_kak_tahun", "")
        values = [number, person.get("nama_personel", ""), person.get("nik", ""),
                  person.get("jabatan_personel", ""), person.get("kualifikasi_pendidikan", ""),
                  person.get("sertifikat_keahlian", ""),
                  f"{min_kak} Tahun" if min_kak not in (None, "") else "",
                  person.get("pengalaman_kerja_bulan", ""), years_display]
        for column, (cell, value) in enumerate(zip(table.add_row().cells, values)):
            align = WD_ALIGN_PARAGRAPH.LEFT if column in (1, 2, 3, 4, 5) else WD_ALIGN_PARAGRAPH.CENTER
            set_cell(cell, value, size=6, align=align)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Render laporan review CV ke DOCX")
    parser.add_argument("--payload-json", required=True, help="JSON payload atau '-' untuk stdin")
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    for stream in (sys.stdin, sys.stdout):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8")
    try:
        payload_text = sys.stdin.read() if args.payload_json == "-" else args.payload_json
        result = render_report(json.loads(payload_text), args.output.resolve())
    except Exception as exc:
        result = {"success": False, "error": type(exc).__name__, "message": str(exc)}
    print("HERMES_CV_REPORT_RESULT=" + json.dumps(result, ensure_ascii=False))
    if not result["success"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
