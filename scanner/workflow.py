"""Ordered Desktop workflow, reusing the existing evidence-review engine.

The XLSX checkpoint is created BEFORE any web dispatch. The final export still
uses review.export so the delivered XLSX and PDF share one verified snapshot.
"""
import hashlib
import json
import sys
import uuid
from pathlib import Path

from . import __version__, review, storage
from .validation import records, source_refs, validate_plan
from .web import normalized

PROTOCOL = 'excel-first-v1'
IDENTITY_FIELDS = ('education', 'certificate_summary', 'claimed_months', 'nik')


def _plan_key(plan):
    data = {key: plan[key] for key in ('roster', 'requirements', 'coverage')}
    return hashlib.sha256(json.dumps(data, sort_keys=True, ensure_ascii=False).encode()).hexdigest()


def require_summary(root, manifest):
    """Legacy 0.4 checkpoints remain readable; new Desktop runs cannot skip XLSX."""
    if manifest.get('workflow') != PROTOCOL:
        return None
    path = root / 'summary.json'
    if not path.is_file():
        raise ValueError('SUMMARY_REQUIRED: buat Excel melalui scanner_review action=summary sebelum verifikasi web/review akhir')
    state = storage.load(path)
    workbook = Path(state['artifact']['path'])
    if (not workbook.is_file() or storage.digest(workbook) != state['artifact']['sha256'] or
            state['plan_key'] != _plan_key(storage.load(root / 'plan.json'))):
        raise ValueError('SUMMARY_STALE: Excel atau pemetaan berubah; periksa checkpoint, jangan lanjut dengan data berbeda')
    return state


def _identity(value):
    if not isinstance(value, dict):
        raise ValueError('identity harus object')
    data = {key: value.get(key) for key in IDENTITY_FIELDS}
    for key in ('education', 'certificate_summary', 'nik'):
        if data[key] is None:
            data[key] = ''
        if not isinstance(data[key], str):
            raise ValueError(f'identity.{key} harus string')
    if data['claimed_months'] is not None and (type(data['claimed_months']) is not int or data['claimed_months'] < 0):
        raise ValueError('claimed_months harus integer nonnegatif atau null; tidak diketahui bukan nol')
    if data['nik']:
        data['nik'] = '[ID_REDACTED]'
    return data


def create_summary(root, manifest, payload):
    from .exporter import personnel_workbook, data_cell, detail_sheet
    from openpyxl import load_workbook

    plan = storage.load(root / 'plan.json')
    if list((root / 'people').glob('*.json')) or list((root / 'web').glob('*.json')):
        existing = require_summary(root, manifest)
        if existing and existing['request'] == payload:
            return {'success': True, 'cached': True, 'artifact': existing['artifact']}
        raise ValueError('Summary tidak boleh diubah setelah verifikasi dimulai; mulai paket baru untuk koreksi pemetaan')
    rows = records(payload.get('people'), 'people')
    roster = {person['id']: person for person in plan['roster']}
    ids = [row.get('id') for row in rows]
    if len(ids) != len(roster) or any(not isinstance(pid, str) for pid in ids) or set(ids) != set(roster):
        raise ValueError('Summary harus memuat tepat SEMUA ID roster, tanpa duplikat')
    mapped = {}
    for row in rows:
        person = roster[row['id']]
        refs = source_refs(root, manifest, row.get('source_refs', []), required=True, kinds={'cv', 'attachment'})
        for ref in refs:
            if not any(c['document_id'] == ref['document_id'] and person['id'] in c.get('person_ids', []) and
                       c['first_page'] <= ref['page'] <= c['last_page'] for c in plan['coverage']):
                raise ValueError('Summary memakai halaman milik personel lain')
        if not any(normalized(person['name']) in normalized(ref['quote']) for ref in person['cv_refs']):
            raise ValueError('Nama personel harus tertulis pada cv_refs; daftar jabatan dalam KAK bukan daftar orang')
        mapped[row['id']] = {'id': row['id'], 'identity': _identity(row.get('identity')), 'source_refs': refs}
    personel, missing_minimum = [], []
    for person in plan['roster']:
        identity = mapped[person['id']]['identity']
        minimums = [r['minimum_months'] for r in plan['requirements'] if r['kind'] == 'experience' and
                    r.get('minimum_months') is not None and normalized(r['role']) in ('*', normalized(person['role']))]
        personel.append({'nama_personel': person['name'], 'nik': identity['nik'], 'jabatan_personel': person['role'],
            'kualifikasi_pendidikan': identity['education'], 'sertifikat_keahlian': identity['certificate_summary'],
            'pengalaman_kerja_bulan': identity['claimed_months'],
            'pengalaman_min_kak_tahun': max(minimums) / 12 if minimums else ''})
        missing_minimum.append(not minimums)
    target = root / 'summary' / uuid.uuid4().hex[:12] / 'Ringkasan_Tenaga_Ahli.xlsx'
    target.parent.mkdir(parents=True)
    book = personnel_workbook({'judul': 'DAFTAR TENAGA AHLI', 'wilayah': manifest['project'],
        'pekerjaan': 'RINGKASAN OCR - belum verifikasi web; angka pengalaman adalah klaim CV', 'personel': personel})
    book.active.title = 'Ringkasan Personel'
    for row, missing in enumerate(missing_minimum, 6):
        if missing:
            data_cell(book.active, row, 7, 'Belum tersedia')
    detail_sheet(book, 'Status Pemeriksaan', ['Item', 'Keterangan'], [
        ['Tahap', 'Ringkasan OCR selesai; verifikasi web dan PDF belum selesai.'],
        ['Run', manifest['run_id']], ['Tanggal acuan', manifest['assessment_date']],
        ['Sumber', 'Data pemetaan dokumen, bukan bukti keabsahan sertifikat.']], [25, 90])
    try:
        book.save(target)
    finally:
        book.close()
    check = load_workbook(target, read_only=True, data_only=True)
    try:
        if check['Ringkasan Personel'].max_row != len(roster) + 5:
            raise ValueError('SUMMARY_PERSONNEL_COUNT_MISMATCH')
    finally:
        check.close()
    artifact = {'name': target.name, 'path': str(target), 'bytes': target.stat().st_size,
                'sha256': storage.digest(target), 'preliminary': True}
    storage.save(root / 'summary.json', {'plan_key': _plan_key(plan), 'people': mapped,
        'artifact': artifact, 'request': payload, 'created_at': storage.now()})
    storage.event(root, 'summary_created', person_count=len(roster))
    return {'success': True, 'artifact': artifact, 'person_count': len(roster)}


def _invalidate_completion(root):
    path = root / 'workflow_result.json'
    if path.exists():
        # Retain exported files and prior results, but do not advertise an old export as current.
        storage.save(path, {'workflow_complete': False})


def _progress(root):
    manifest = storage.load(root / 'manifest.json')
    pending = []
    for doc in manifest['documents']:
        path = root / 'ocr' / f"{doc['id']}.json"
        if not path.exists() or storage.load(path).get('processing'):
            pending.append(doc['id'])
    if pending:
        return {'stage': 'ocr', 'pending_document_ids': pending, 'next_action': 'document'}
    if not (root / 'plan.json').exists():
        return {'stage': 'mapping', 'next_action': 'read_document kemudian plan'}
    if manifest.get('workflow') == PROTOCOL and not (root / 'summary.json').exists():
        return {'stage': 'summary', 'next_action': 'summary'}
    require_summary(root, manifest)
    if manifest['allow_web'] and not list((root / 'web').glob('*.json')):
        return {'stage': 'web', 'next_action': 'scanner_web_lookup'}
    plan = storage.load(root / 'plan.json')
    pending = [p['id'] for p in plan['roster'] if not (root / 'people' / f"{p['id']}.json").exists()]
    if pending:
        return {'stage': 'analysis', 'pending_person_ids': pending, 'next_action': 'save_person'}
    return {'stage': 'export', 'next_action': 'export'}


def dispatch(args):
    action = args.get('action')
    payload = args.get('payload', {})
    if not isinstance(payload, dict):
        raise ValueError('payload harus object')
    if action == 'health':
        expected = args.get('_plugin_version', __version__)
        result = {'success': expected == __version__, 'plugin_version': expected, 'runtime_version': __version__,
                  'protocol': PROTOCOL, 'runtime_root': str(Path(__file__).resolve().parents[1]),
                  'workflow_complete': False, 'next_action': 'help kemudian start', 'gpu_tested': False}
        if not result['success']:
            result.update(error='PLUGIN_RUNTIME_MISMATCH', message='Sinkronkan plugin dan runtime, restart Hermes. Jangan membuat shim atau mengedit kode saat scan.')
        return result
    if action == 'help':
        root = Path(__file__).resolve().parents[1]
        return {'success': True, 'protocol': PROTOCOL, 'workflow_complete': False,
                'format': (root / 'docs' / 'WORKFLOW.md').read_text(encoding='utf-8') + '\n\n' +
                          (root / 'docs' / 'REVIEW_FORMAT.md').read_text(encoding='utf-8'),
                'next_action': 'start'}
    if action == 'start':
        if args.get('_plugin_version', __version__) != __version__:
            raise ValueError('PLUGIN_RUNTIME_MISMATCH: sinkronkan plugin dan restart Hermes')
        if payload.get('workflow_version', __version__) != __version__:
            raise ValueError('DESKTOP_VERSION_MISMATCH: reload plugin Desktop sebelum mulai paket')
        result = review.start(payload)
        with storage.locked(result['run_id']) as root:
            manifest = storage.load(root / 'manifest.json')
            manifest['workflow'] = PROTOCOL
            storage.save(root / 'manifest.json', manifest)
        return {**result, 'protocol': PROTOCOL, 'stage': 'ocr', 'workflow_complete': False,
                'next_action': 'document untuk SETIAP ID terpilih; bukan pencarian file lain'}
    run_id = args.get('run_id')
    if action in {'summary', 'verify_kak', 'status'}:
        with storage.locked(run_id) as root:
            manifest = storage.load(root / 'manifest.json')
            if action == 'summary':
                result = create_summary(root, manifest, payload)
                _invalidate_completion(root)
                return {**result, 'stage': 'summary_done', 'workflow_complete': False,
                        'next_action': 'scanner_web_lookup lalu verify_kak/save_person' if manifest['allow_web'] else 'save_person dengan keterbatasan web',
                        'instruction': 'Excel ringkasan sudah dibuat. JANGAN berhenti di sini. Lanjutkan validasi dan PDF; kirim hanya dua file final dari export.'}
            if action == 'verify_kak':
                require_summary(root, manifest)
                if set(payload) - {'analysis', 'receipt_ids'}:
                    raise ValueError('verify_kak hanya memperbarui analysis dan receipt_ids; tidak mengubah persyaratan atau paket')
                plan = storage.load(root / 'plan.json')
                plan['kak'].update(payload)
                plan = validate_plan(root, manifest, plan)
                storage.save(root / 'plan.json', plan)
                _invalidate_completion(root)
                return {'success': True, 'workflow_complete': False, 'kak_status': plan['kak']['status'], 'next_action': 'save_person untuk seluruh roster'}
            done_path = root / 'workflow_result.json'
            done = storage.load(done_path) if done_path.exists() else {}
            if done.get('workflow_complete') and all(Path(a['path']).is_file() and storage.digest(a['path']) == a['sha256'] for a in done.get('artifacts', [])):
                return {**done, 'stage': 'complete', 'next_action': 'kirim chat_markdown'}
            return {'success': True, 'run_id': run_id, 'workflow_complete': False, **_progress(root)}
    with storage.locked(run_id) as root:
        manifest = storage.load(root / 'manifest.json')
        if action == 'plan' and (root / 'summary.json').exists():
            raise ValueError('Pemetaan sudah menjadi Excel. Gunakan verify_kak untuk receipt/analisis KAK, atau run baru untuk mengubah roster/kriteria')
        if action in {'save_person', 'export'}:
            summary = require_summary(root, manifest)
            if manifest.get('workflow') == PROTOCOL and manifest['allow_web'] and not list((root / 'web').glob('*.json')):
                raise ValueError('WEB_REVIEW_REQUIRED: jalankan scanner_web_lookup setelah Excel; kegagalan web juga harus tercatat sebagai receipt')
            if action == 'save_person' and summary:
                expected = summary['people'].get(payload.get('id'))
                if not expected or _identity(payload.get('identity')) != expected['identity']:
                    raise ValueError('SUMMARY_IDENTITY_MISMATCH: gunakan identity dari Excel pemetaan; jangan mengubah biodata/angka tanpa run pemetaan baru')
        if action in {'plan', 'save_person'}:
            _invalidate_completion(root)
    # The original review engine owns its locks, validation, OCR and renderers.
    result = review.dispatch(args)
    if action == 'export':
        result['workflow_complete'] = bool(result['success'] and len(result.get('artifacts', [])) == 2)
        result['stage'] = 'complete' if result['workflow_complete'] else 'export_partial'
        result['next_action'] = 'kirim chat_markdown' if result['workflow_complete'] else 'laporkan kegagalan dan hanya file yang berhasil'
        with storage.locked(run_id) as root:
            storage.save(root / 'workflow_result.json', result)
        return result
    return {**result, 'workflow_complete': False,
            'next_action': {'plan': 'summary', 'save_person': 'status; lanjut personel lain atau export',
                            'document': 'read_document; selesaikan semua dokumen lalu plan'}.get(action, 'ikuti next dari hasil tool'),
            'instruction': result.get('instruction', '') + ' success hanya berarti tahap ini berhasil. Jangan menyatakan scan lengkap sebelum workflow_complete=true dan ada dua artefak final.'}


def main():
    for stream in (sys.stdin, sys.stdout):
        if hasattr(stream, 'reconfigure'):
            stream.reconfigure(encoding='utf-8')
    try:
        result = dispatch(json.load(sys.stdin))
    except Exception as exc:
        result = {'success': False, 'workflow_complete': False, 'error': type(exc).__name__, 'message': str(exc)}
    print('HERMES_REVIEW_RESULT=' + json.dumps(result, ensure_ascii=False), flush=True)
    if not result['success']:
        raise SystemExit(1)


if __name__ == '__main__':
    main()
