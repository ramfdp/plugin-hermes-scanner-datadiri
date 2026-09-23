"""Legacy personnel XLSX exporter; also used by the review bundle renderer."""
import argparse
import json
import math
import sys
import textwrap
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins

TEMPLATE_ID = "daftar_tenaga_ahli"
REQUIRED_FIELDS = ("judul", "wilayah", "pekerjaan", "personel")
HEADERS = ["No", "Nama Personel", "NIK", "Jabatan Personel", "Kualifikasi\nPendidikan",
           "Sertifikat Keahlian", "Pengalaman min\ndalam KAK\n(Tahun)",
           "Pengalaman\nKerja (Bulan)", "Pengalaman\nKerja (Tahun)"]
HISTORY_FIELDS = ("nama_personel", "employer", "role", "start_date", "end_date", "duration_months",
                  "responsibilities", "project", "source_page", "source_quote")


def load_json(path):
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"File JSON tidak ditemukan: {path}")
    if not path.stat().st_size:
        raise ValueError(f"File JSON kosong: {path}")
    try:
        return json.loads(path.read_text(encoding="utf-8-sig"))
    except json.JSONDecodeError as exc:
        raise ValueError(f"JSON tidak valid pada '{path}' baris {exc.lineno}, kolom {exc.colno}: {exc.msg}") from exc


def truncate_decimal(value, decimals=2):
    multiplier = 10 ** decimals
    return math.floor(value * multiplier) / multiplier


def calculate_years(months):
    try:
        return truncate_decimal(float(months) / 12, 2)
    except (TypeError, ValueError):
        return 0.0


def format_indonesian_decimal(value):
    try:
        return f"{float(value):.2f}".replace(".", ",")
    except (TypeError, ValueError):
        return ""


def validate_payload(payload, schema=None):
    if not isinstance(payload, dict):
        raise ValueError("Payload harus berupa object.")
    required = REQUIRED_FIELDS if schema is None else schema.get("required_fields", [])
    missing = [field for field in required if field not in payload]
    if missing:
        raise ValueError("Field wajib belum tersedia: " + ", ".join(missing))
    for key in ("personel", "employment_history"):
        rows = payload.get(key, [])
        if not isinstance(rows, list) or any(not isinstance(row, dict) for row in rows):
            raise ValueError(f"'{key}' harus berupa array object.")


def data_cell(sheet, row, col, value):
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False)
    if isinstance(value, str) and len(value) > 32767:
        raise ValueError("Teks melampaui batas sel Excel; pisahkan detail menjadi beberapa baris")
    cell = sheet.cell(row, col, value)
    if isinstance(value, str):
        cell.data_type = "s"  # OCR and web content must never become spreadsheet formulas.
    return cell


def fit_row(sheet, row, values, widths, *, minimum=28, size=10):
    lines = max((sum(max(1, len(textwrap.wrap(str(line), max(4, int(width)-2))))
                      for line in str(value if value is not None else "").split('\n'))
                 for value, width in zip(values, widths)), default=1)
    sheet.row_dimensions[row].height = min(409, max(minimum, lines * (size + 4) + 10))


def detail_sheet(workbook, name, headers, rows, widths=None):
    sheet = workbook.create_sheet(name)
    sheet.sheet_view.showGridLines = False
    widths = widths or [28] * len(headers)
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = width
    for row_index, values in enumerate([headers, *rows], 1):
        for col, value in enumerate(values, 1):
            cell = data_cell(sheet, row_index, col, value)
            cell.font = Font(name="Calibri", size=10, bold=row_index == 1, color="FFFFFF" if row_index == 1 else "17293D")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.fill = PatternFill("solid", fgColor="264560" if row_index == 1 else ("F1F5F9" if row_index % 2 else "FFFFFF"))
        fit_row(sheet, row_index, values, widths, minimum=32)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.print_title_rows = "1:1"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    return sheet


def personnel_workbook(payload):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Daftar Tenaga Ahli"
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_margins = PageMargins(left=.25, right=.25, top=.35, bottom=.35, header=.15, footer=.15)
    centered = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    border = Border(**{key: Side(style="thin", color="526579") for key in ("left", "right", "top", "bottom")})
    titles = [payload.get("judul", "DAFTAR TENAGA AHLI"), payload.get("wilayah", ""), payload.get("pekerjaan", "")]
    for row, value in enumerate(titles, 1):
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
        cell = data_cell(sheet, row, 1, value)
        cell.font = Font(name="Century Gothic", size=11 if row == 1 else 10, bold=True)
        cell.alignment = centered
        sheet.row_dimensions[row].height = 24
    sheet.row_dimensions[4].height = 10
    for col, value in enumerate(HEADERS, 1):
        cell = data_cell(sheet, 5, col, value)
        cell.font = Font(name="Century Gothic", size=8, bold=True)
        cell.fill = PatternFill("solid", fgColor="8EA9D8")
        cell.alignment, cell.border = centered, border
    sheet.row_dimensions[5].height = 48
    widths = (5, 24, 24, 36, 28, 34, 19, 17, 17)
    for col, width in zip("ABCDEFGHI", widths):
        sheet.column_dimensions[col].width = width
    for number, person in enumerate(payload.get("personel", []), 1):
        months = person.get("pengalaman_kerja_bulan", "")
        years = person.get("pengalaman_kerja_tahun")
        if years in (None, ""):
            years = calculate_years(months) if months not in (None, "") else ""
        minimum = person.get("pengalaman_min_kak_tahun", "")
        values = [number, person.get("nama_personel", ""), str(person.get("nik", "")),
                  person.get("jabatan_personel", ""), person.get("kualifikasi_pendidikan", ""),
                  person.get("sertifikat_keahlian", ""), f"{minimum} Tahun" if minimum not in (None, "") else "",
                  months, format_indonesian_decimal(years) if years != "" else ""]
        for col, value in enumerate(values, 1):
            cell = data_cell(sheet, number + 5, col, value)
            cell.font, cell.border = Font(name="Century Gothic", size=8), border
            cell.alignment = left if col in (2, 3) else centered
        sheet.cell(number + 5, 3).number_format = "@"
        sheet.cell(number + 5, 8).number_format = "0"
        fit_row(sheet, number + 5, values, widths, minimum=54, size=8)
    last = max(5, 5 + len(payload.get("personel", [])))
    sheet.print_area = f"A1:I{last}"
    sheet.auto_filter.ref = f"A5:I{last}"
    sheet.print_title_rows = "1:5"
    sheet.freeze_panes = "A6"
    sheet.print_options.horizontalCentered = True
    if "employment_history" in payload:
        detail_sheet(workbook, "Riwayat Pekerjaan", list(HISTORY_FIELDS),
                     [[job.get(key, "") for key in HISTORY_FIELDS] for job in payload["employment_history"]])
    return workbook


def render_daftar_tenaga_ahli(payload, output_path):
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    book = personnel_workbook(payload)
    try:
        book.save(output_path)
    finally:
        book.close()


def export_payload(template_id, payload, output_path):
    if template_id != TEMPLATE_ID:
        raise ValueError(f"Renderer untuk template '{template_id}' belum tersedia.")
    validate_payload(payload)
    output_path = Path(output_path)
    render_daftar_tenaga_ahli(payload, output_path)
    from .storage import save
    payload_file = output_path.with_suffix(".payload.json")
    save(payload_file, payload)
    book = load_workbook(output_path, read_only=True, data_only=True)
    try:
        data = {s.title: list(s.iter_rows(min_row=5 if s.title == "Daftar Tenaga Ahli" else 1, values_only=True)) for s in book}
    finally:
        book.close()
    return {"success": True, "template": template_id, "output_file": str(output_path.resolve()),
            "row_count": len(payload["personel"]), "payload_file": str(payload_file.resolve()), "workbook_data": data}


def export_template(template_id, payload_path, output_path):
    result = export_payload(template_id, load_json(payload_path), output_path)
    result["payload_file"] = str(Path(payload_path).resolve())
    return result


def main():
    parser = argparse.ArgumentParser(description="Personnel XLSX exporter")
    parser.add_argument("--template", required=True)
    group = parser.add_mutually_exclusive_group(required=True)
    group.add_argument("--payload", type=Path)
    group.add_argument("--payload-json")
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    for stream in (sys.stdin, sys.stdout):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8")
    try:
        result = (export_template(args.template, args.payload, args.output) if args.payload_json is None else
                  export_payload(args.template, json.loads(sys.stdin.read() if args.payload_json == "-" else args.payload_json), args.output))
    except Exception as exc:
        result = {"success": False, "error": type(exc).__name__, "message": str(exc)}
    print("HERMES_EXPORT_RESULT=" + json.dumps(result, ensure_ascii=False))
    if not result["success"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
