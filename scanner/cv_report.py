"""Compatibility DOCX renderer; new multi-CV reviews use renderers/pdf.py."""
import argparse
import json
import sys
from pathlib import Path
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from .storage import save


def text(value):
    return "" if value is None else json.dumps(value, ensure_ascii=False) if isinstance(value, (list, dict)) else str(value)


def _add_table(document, rows):
    table = document.add_table(rows=0, cols=len(rows[0]))
    table.style = "Table Grid"
    for row in rows:
        for cell, value in zip(table.add_row().cells, row):
            cell.text = text(value)
    return table


def render_report(payload, output_path):
    required = ("cv_file", "biodata", "experience_validation", "attachment_cross_check", "findings")
    if not isinstance(payload, dict):
        raise ValueError("Payload harus berupa object.")
    missing = [key for key in required if key not in payload]
    if missing:
        raise ValueError("Field wajib belum tersedia: " + ", ".join(missing))
    for key in ("experience_validation", "attachment_cross_check", "employer_validation"):
        rows = payload.get(key) or []
        if not isinstance(rows, list) or any(not isinstance(row, dict) for row in rows):
            raise ValueError(f"'{key}' harus berupa array object.")
    for key in ("findings", "internet_sources", "chronology_validation"):
        rows = payload.get(key)
        if rows is not None and (not isinstance(rows, list) or any(not isinstance(row, str) for row in rows)):
            raise ValueError(f"'{key}' harus berupa array string.")
    if not isinstance(payload.get("source_artifacts", {}), dict):
        raise ValueError("source_artifacts harus berupa object")
    for employer in payload.get("employer_validation", []):
        sources = employer.get("sources", [])
        if not isinstance(sources, list) or any(not isinstance(url, str) for url in sources):
            raise ValueError("sources harus berupa array URL")
        if employer.get("status") in ("terverifikasi", "sebagian terverifikasi") and not any(url.startswith(("https://", "http://")) for url in sources):
            raise ValueError("Status perusahaan terverifikasi memerlukan URL sumber")
    doc = Document()
    doc.add_paragraph("LAPORAN REVIEW CV", "Title")
    doc.add_paragraph("Dokumen hasil pemetaan biodata, validasi pengalaman, dan cross-check lampiran.")
    doc.add_paragraph(f"File CV: {payload['cv_file']}")
    if payload.get("kak_file"):
        doc.add_paragraph(f"File KAK: {payload['kak_file']}")
    doc.add_heading("1. Pemetaan Biodata", 1)
    biodata = payload["biodata"] if isinstance(payload["biodata"], dict) else {"biodata": payload["biodata"]}
    _add_table(doc, [["Field", "Nilai"], *[[key.replace('_', ' ').title(), value] for key, value in biodata.items()]])
    doc.add_heading("2. Kesesuaian Pengalaman dengan Posisi dalam KAK", 1)
    fields = ("kak_requirement", "cv_claim", "internet_validation", "status", "notes")
    rows = [["Posisi/Kriteria KAK", "Klaim CV", "Validasi Internet", "Status", "Catatan"]]
    rows += [[item.get(key, "") for key in fields] for item in payload["experience_validation"] or []]
    _add_table(doc, rows if len(rows) > 1 else [["Tidak ada pengalaman yang dapat divalidasi."]])
    doc.add_heading("3. Cross-check CV dan Lampiran", 1)
    fields = ("attachment", "field", "cv_value", "attachment_value", "status")
    rows = [["Lampiran", "Field", "CV", "Lampiran", "Status/Catatan"]]
    rows += [[item.get(key, "") for key in fields] for item in payload["attachment_cross_check"] or []]
    _add_table(doc, rows if len(rows) > 1 else [["Tidak ada lampiran yang dicross-check."]])
    for title, values in (("4. Temuan/Janggal yang Perlu Dikomentari", payload["findings"] or ["Tidak ada temuan."]),
                          ("5. Sumber Validasi Internet", payload.get("internet_sources") or ["Tidak ada sumber yang dicatat."])):
        doc.add_heading(title, 1)
        for value in values:
            doc.add_paragraph(value, "List Bullet")
    doc.add_heading("6. Kesimpulan", 1)
    doc.add_paragraph(text(payload.get("conclusion", "Belum ada kesimpulan.")))
    doc.add_paragraph("Status akhir: " + text(payload.get("status", "perlu review manual")))
    if "employer_validation" in payload:
        doc.add_heading("7. Verifikasi Perusahaan", 1)
        _add_table(doc, [["Perusahaan", "Status", "Bukti", "Sumber", "Diperiksa"], *[
            [item.get("employer"), item.get("status"), item.get("evidence"), '\n'.join(item.get("sources", [])), item.get("checked_at")]
            for item in payload["employer_validation"]]])
        doc.add_paragraph("Keberadaan perusahaan bukan bukti kandidat bekerja di sana.")
    if "chronology_validation" in payload:
        doc.add_heading("8. Validasi Kronologi", 1)
        for value in payload["chronology_validation"]:
            doc.add_paragraph(value)
    if payload.get("source_artifacts"):
        doc.add_heading("9. Dokumen Sumber", 1)
        _add_table(doc, [["Artefak", "Lokasi"], *payload["source_artifacts"].items()])
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    json_file = output_path.with_suffix(".json")
    save(json_file, payload)
    return {"success": True, "template": "cv_review", "output_file": str(output_path.resolve()),
            "json_file": str(json_file.resolve()), "finding_count": len(payload["findings"] or [])}


def render_personnel_docx(payload, output_path):
    from .exporter import HEADERS
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = section.bottom_margin = Inches(.35)
    section.left_margin = section.right_margin = Inches(.25)
    for key, default in (("judul", "DAFTAR TENAGA AHLI"), ("wilayah", ""), ("pekerjaan", "")):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text(payload.get(key, default)))
        r.bold, r.font.name, r.font.size = True, "Century Gothic", Pt(11)
    rows = [HEADERS]
    for number, item in enumerate(payload["personel"], 1):
        minimum, years = item.get("pengalaman_min_kak_tahun", ""), item.get("pengalaman_kerja_tahun", "")
        rows.append([number, item.get("nama_personel"), item.get("nik"), item.get("jabatan_personel"),
                     item.get("kualifikasi_pendidikan"), item.get("sertifikat_keahlian"),
                     f"{minimum} Tahun" if minimum not in (None, "") else "", item.get("pengalaman_kerja_bulan"),
                     "" if years in (None, "") else f"{float(years):.2f}".replace('.', ',')])
    table = _add_table(doc, rows)
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.name, r.font.size = "Century Gothic", Pt(7)
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--payload-json", required=True)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    for stream in (sys.stdin, sys.stdout):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8")
    try:
        result = render_report(json.loads(sys.stdin.read() if args.payload_json == "-" else args.payload_json), args.output)
    except Exception as exc:
        result = {"success": False, "error": type(exc).__name__, "message": str(exc)}
    print("HERMES_CV_REPORT_RESULT=" + json.dumps(result, ensure_ascii=False))
    if not result["success"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
