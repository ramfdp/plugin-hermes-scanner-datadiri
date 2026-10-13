"""Final review regressions for Scanner Data Diri 0.7.0."""
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from scanner import review, storage
from scanner.final_review import evaluate_anomalies, evaluate_cv_supporting_documents
from helpers import create_run, plan_payload, person_payload


class FinalReviewV070Test(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.root = Path(temp.name)
        env = patch.dict(os.environ, HERMES_SCANNER_OUTPUT_DIR=str(self.root / 'output'))
        env.start()
        self.addCleanup(env.stop)

    @staticmethod
    def ref(page, quote):
        return {'document_id': 'D001', 'page': page, 'quote': quote}

    def prepare_run(self, support_project='Proyek Jembatan Sungai'):
        source = self.root / 'final-review.pdf'
        source.write_bytes(b'Synthetic PDF; OCR mocked')
        started = review.dispatch({'action': 'start', 'payload': {
            'project': 'Paket Final Review Sintetis',
            'assessment_date': '2026-09-30',
            'allow_web': False,
            'documents': [{'kind': 'cv', 'path': str(source)}],
        }})
        pages = [
            'Personel Final\nS1 Teknik Sipil\nSite Engineer\nProyek Jembatan Sungai\n'
            'Konsultan: PT Konsultan A | Pemberi Kerja: Dinas PU\n'
            'Kontraktor: PT Kontraktor A\nPengawasan struktur jembatan\n2020-01 sampai 2021-12',
            'IJAZAH\nPersonel Final\nIjazah S1 Teknik Sipil\nSarjana Teknik\nUniversitas Contoh',
            'CERT-001 Personel Final Ahli Teknik Jembatan BNSP Level 7',
            'KAK POSISI ENGINEER\n'
            'Posisi Engineer membutuhkan pengalaman sebagai Site Engineer\n'
            'Pengalaman proyek jembatan\n'
            'Pengalaman pada konsultan supervisi\n'
            'Pengalaman minimal 12 bulan\n'
            'Uraian tugas pengawasan struktur\n'
            'Pendidikan minimal S1 Teknik Sipil\n'
            'SKK Ahli Teknik Jembatan jenjang 7 masih berlaku\n'
            'Kompetensi teknis pengawasan struktur',
            f'SURAT PENGALAMAN\n{support_project}\nPT Konsultan A\nSite Engineer\n'
            'Pemberi Kerja Dinas PU\nKonsultan PT Konsultan A\nKontraktor PT Kontraktor A\n'
            '2020-01 sampai 2021-12',
        ]

        def extract(path, output, model):
            self.assertTrue(path.samefile(source))
            for index, text in enumerate(pages, 1):
                (output / f'page_{index:04d}.md').write_text(text, encoding='utf-8')
                (output / f'page_{index:04d}.json').write_text('{}', encoding='utf-8')
            return len(pages)

        with patch('scanner.ocr.run_mineru', side_effect=extract):
            result = review.dispatch({'action': 'document', 'run_id': started['run_id'],
                                      'payload': {'document_id': 'D001'}})
        self.assertTrue(result['success'], result)

        plan = {
            'kak': {'title': 'KAK Posisi Engineer', 'version': '1', 'package_id': 'PKT-070',
                    'analysis': 'KAK sintetis memuat delapan requirement structured matching untuk final review.',
                    'receipt_ids': []},
            'requirements': [
                {'id': 'R1', 'role': 'Engineer', 'kind': 'experience',
                 'audit_code': 'position_experience_match',
                 'text': 'Posisi Engineer membutuhkan pengalaman sebagai Site Engineer',
                 'minimum_months': None,
                 'source_refs': [self.ref(4, 'Posisi Engineer membutuhkan pengalaman sebagai Site Engineer')]},
                {'id': 'R2', 'role': 'Engineer', 'kind': 'experience',
                 'audit_code': 'project_kak_match',
                 'text': 'Pengalaman proyek jembatan', 'minimum_months': None,
                 'source_refs': [self.ref(4, 'Pengalaman proyek jembatan')]},
                {'id': 'R3', 'role': 'Engineer', 'kind': 'experience',
                 'audit_code': 'organization_role_match',
                 'text': 'Pengalaman pada konsultan supervisi', 'minimum_months': None,
                 'source_refs': [self.ref(4, 'Pengalaman pada konsultan supervisi')]},
                {'id': 'R5', 'role': 'Engineer', 'kind': 'experience',
                 'audit_code': 'experience_duration_match',
                 'text': 'Pengalaman minimal 12 bulan', 'minimum_months': 12,
                 'source_refs': [self.ref(4, 'Pengalaman minimal 12 bulan')]},
                {'id': 'R6', 'role': 'Engineer', 'kind': 'experience',
                 'audit_code': 'responsibility_position_match',
                 'text': 'Uraian tugas pengawasan struktur', 'minimum_months': None,
                 'source_refs': [self.ref(4, 'Uraian tugas pengawasan struktur')]},
                {'id': 'R7', 'role': 'Engineer', 'kind': 'education',
                 'audit_code': 'education_major_match',
                 'text': 'Pendidikan minimal S1 Teknik Sipil', 'minimum_months': None,
                 'parameters': {'minimum_level': 'S1', 'accepted_majors': ['Teknik Sipil']},
                 'source_refs': [self.ref(4, 'Pendidikan minimal S1 Teknik Sipil')]},
                {'id': 'R8', 'role': 'Engineer', 'kind': 'certificate',
                 'audit_code': 'certificate_kak_validity',
                 'text': 'SKK Ahli Teknik Jembatan jenjang 7 masih berlaku', 'minimum_months': None,
                 'parameters': {'schemes': ['Ahli Teknik Jembatan'], 'levels': ['7'],
                                'issuers': [], 'require_current': True},
                 'source_refs': [self.ref(4, 'SKK Ahli Teknik Jembatan jenjang 7 masih berlaku')]},
                {'id': 'R12', 'role': 'Engineer', 'kind': 'other',
                 'audit_code': 'technical_competency_match',
                 'text': 'Kompetensi teknis pengawasan struktur', 'minimum_months': None,
                 'source_refs': [self.ref(4, 'Kompetensi teknis pengawasan struktur')]},
            ],
            'roster': [{'id': 'P1', 'name': 'Personel Final', 'role': 'Engineer',
                        'cv_refs': [self.ref(1, 'Personel Final')],
                        'certificate_inventory': [{'id': 'C1', 'source_refs': [
                            self.ref(3, 'CERT-001 Personel Final Ahli Teknik Jembatan BNSP Level 7')]}]}],
            'coverage': [
                {'document_id': 'D001', 'first_page': 1, 'last_page': 3, 'person_ids': ['P1']},
                {'document_id': 'D001', 'first_page': 4, 'last_page': 4,
                 'person_ids': [], 'reason': 'reference'},
                {'document_id': 'D001', 'first_page': 5, 'last_page': 5, 'person_ids': ['P1']},
            ],
            'page_classification': [
                {'document_id': 'D001', 'first_page': 1, 'last_page': 1, 'kind': 'cv'},
                {'document_id': 'D001', 'first_page': 2, 'last_page': 3, 'kind': 'attachment'},
                {'document_id': 'D001', 'first_page': 4, 'last_page': 4, 'kind': 'kak'},
                {'document_id': 'D001', 'first_page': 5, 'last_page': 5, 'kind': 'attachment'},
            ],
        }
        review.dispatch({'action': 'plan', 'run_id': started['run_id'], 'payload': plan})
        return started['run_id']

    def person(self, support_project='Proyek Jembatan Sungai', expires_on='2030-12-31'):
        semantic = [
            ('R1', 'Site Engineer'),
            ('R2', 'Proyek Jembatan Sungai'),
            ('R3', 'Konsultan: PT Konsultan A | Pemberi Kerja: Dinas PU'),
            ('R6', 'Pengawasan struktur jembatan'),
            ('R12', 'Pengawasan struktur jembatan'),
        ]
        return {
            'id': 'P1',
            'identity': {'education': 'S1 Teknik Sipil',
                         'certificate_summary': 'Ahli Teknik Jembatan Level 7',
                         'claimed_months': 24, 'nik': ''},
            'summary': 'Fixture final review memuat CV, pendidikan, sertifikat, pengalaman pendukung dan KAK terstruktur.',
            'checks': [],
            'education_records': [{
                'level': 'S1', 'degree': 'Sarjana Teknik', 'major': 'Teknik Sipil',
                'institution': 'Universitas Contoh',
                'source_refs': [self.ref(1, 'S1 Teknik Sipil')],
                'supporting_refs': [
                    self.ref(2, 'Ijazah S1 Teknik Sipil'),
                    self.ref(2, 'Sarjana Teknik'),
                    self.ref(2, 'Universitas Contoh'),
                ],
            }],
            'employment_history': [{
                'employer': 'PT Konsultan A', 'role': 'Site Engineer',
                'client': 'Dinas PU', 'consultant': 'PT Konsultan A',
                'contractor': 'PT Kontraktor A', 'represented_organization': '',
                'project': 'Proyek Jembatan Sungai',
                'start_date': '2020-01', 'end_date': '2021-12',
                'relevant': False, 'responsibilities': 'Pengawasan struktur jembatan',
                'source_refs': [
                    self.ref(1, 'Site Engineer'),
                    self.ref(1, 'Proyek Jembatan Sungai'),
                    self.ref(1, 'Konsultan: PT Konsultan A | Pemberi Kerja: Dinas PU'),
                    self.ref(1, 'Pengawasan struktur jembatan'),
                ],
                'supporting_refs': [self.ref(5, 'SURAT PENGALAMAN')],
                'supporting_facts': [{
                    'project': support_project, 'employer': 'PT Konsultan A',
                    'role': 'Site Engineer', 'client': 'Dinas PU',
                    'consultant': 'PT Konsultan A', 'contractor': 'PT Kontraktor A',
                    'strict_evidence': True,
                    'start_date': '2020-01', 'end_date': '2021-12',
                    'source_refs': [
                        self.ref(5, support_project),
                        self.ref(5, 'PT Konsultan A'),
                        self.ref(5, 'Site Engineer'),
                        self.ref(5, 'Dinas PU'),
                        self.ref(5, 'PT Kontraktor A'),
                        self.ref(5, '2020-01 sampai 2021-12'),
                    ],
                }],
            }],
            'history_extraction': {
                'status': 'complete', 'identified_count': 1,
                'analysis': 'Satu pengalaman yang dapat diidentifikasi pada CV sudah dipetakan seluruhnya.',
                'source_refs': [self.ref(1, 'Proyek Jembatan Sungai')],
            },
            'semantic_assessments': [{
                'requirement_id': requirement_id, 'status': 'memenuhi',
                'finding': f'Bukti kandidat mendukung requirement {requirement_id}.',
                'analysis': 'Assessment semantic sintetis membandingkan requirement atomik dengan evidence kandidat.',
                'experience_ids': ['E001'],
                'source_refs': [self.ref(1, quote)],
            } for requirement_id, quote in semantic],
            'employer_checks': [{
                'employer': 'PT Konsultan A',
                'analysis': 'Fixture final review tidak melakukan verifikasi web perusahaan.',
            }],
            'certificates': [{
                'inventory_id': 'C1', 'number': 'CERT-001', 'holder': 'Personel Final',
                'scheme': 'Ahli Teknik Jembatan', 'issuer': 'BNSP', 'level': '7',
                'issued_on': '2025-01-01', 'expires_on': expires_on,
                'analysis': 'Sertifikat sintetis memuat skema dan jenjang yang dapat dibandingkan dengan requirement KAK.',
                'source_refs': [self.ref(3, 'CERT-001 Personel Final Ahli Teknik Jembatan BNSP Level 7')],
            }],
            'certificate_limitation': '',
            'findings': [],
            'audit_checks': [],
        }

    def test_all_fifteen_audits_end_in_memenuhi_when_evidence_is_consistent(self):
        rid = self.prepare_run()
        result = review.dispatch({'action': 'save_person', 'run_id': rid, 'payload': self.person()})
        statuses = {row['code']: row['status'] for row in result['audit_checks']}
        self.assertEqual(len(statuses), 15)
        self.assertEqual(statuses['cv_supporting_document_match'], 'memenuhi')
        self.assertEqual(statuses['data_anomaly'], 'memenuhi')
        self.assertEqual(statuses['kak_conclusion'], 'memenuhi')
        self.assertEqual(result['final_conclusion']['label'], 'Memenuhi')
        self.assertEqual(result['anomaly_count'], 0)

        stored = storage.load(storage.run_dir(rid) / 'people/P1.json')
        self.assertEqual(stored['document_cross_check']['conflicts'], [])
        self.assertEqual(stored['document_cross_check']['missing'], [])
        self.assertEqual(stored['anomalies'], [])
        self.assertEqual(stored['final_conclusion']['status'], 'memenuhi')

    def test_cross_document_conflict_becomes_clarification_and_anomaly(self):
        support_project = 'Proyek Jalan Raya'
        rid = self.prepare_run(support_project=support_project)
        result = review.dispatch({'action': 'save_person', 'run_id': rid,
                                  'payload': self.person(support_project=support_project)})
        statuses = {row['code']: row['status'] for row in result['audit_checks']}
        self.assertEqual(statuses['cv_supporting_document_match'], 'perlu_klarifikasi')
        self.assertEqual(statuses['data_anomaly'], 'perlu_klarifikasi')
        self.assertEqual(statuses['kak_conclusion'], 'perlu_klarifikasi')
        self.assertEqual(result['final_conclusion']['label'], 'Perlu Klarifikasi')
        self.assertGreater(result['anomaly_count'], 0)

        stored = storage.load(storage.run_dir(rid) / 'people/P1.json')
        self.assertTrue(any(item['code'] == 'cv_support_conflict' for item in stored['anomalies']))
        self.assertEqual([item['id'] for item in stored['anomalies']],
                         [f"A{i:03d}" for i in range(1, len(stored['anomalies']) + 1)])

    def test_kak_failure_overrides_other_successes(self):
        rid = self.prepare_run()
        result = review.dispatch({'action': 'save_person', 'run_id': rid,
                                  'payload': self.person(expires_on='2025-01-01')})
        statuses = {row['code']: row['status'] for row in result['audit_checks']}
        self.assertEqual(statuses['certificate_kak_validity'], 'tidak_memenuhi')
        self.assertEqual(statuses['kak_conclusion'], 'tidak_memenuhi')
        self.assertEqual(result['final_conclusion']['label'], 'Tidak Memenuhi')

    def test_no_kak_is_never_memenuhi(self):
        rid = create_run(self.root)
        plan = plan_payload()
        plan['requirements'] = []
        plan['kak']['analysis'] = 'Tidak ada persyaratan KAK yang digunakan pada pengujian final conclusion ini.'
        review.dispatch({'action': 'plan', 'run_id': rid, 'payload': plan})
        person = person_payload(0)
        person['checks'] = []
        result = review.dispatch({'action': 'save_person', 'run_id': rid, 'payload': person})
        self.assertEqual(result['final_conclusion']['status'], 'perlu_klarifikasi')
        self.assertEqual(result['final_conclusion']['label'], 'Perlu Klarifikasi')

    def test_payload_cannot_spoof_final_audits(self):
        rid = self.prepare_run()
        person = self.person()
        person['audit_checks'] = [{
            'code': 'kak_conclusion', 'status': 'memenuhi',
            'finding': 'Payload mencoba memaksa final conclusion.',
            'analysis': 'Kesimpulan final tidak boleh berasal dari payload model.',
            'source_refs': [self.ref(1, 'Personel Final')], 'kak_refs': [],
        }]
        with self.assertRaisesRegex(ValueError, 'server-side'):
            review.dispatch({'action': 'save_person', 'run_id': rid, 'payload': person})

    def test_cross_check_anomaly_keeps_only_record_evidence(self):
        first = self.ref(1, 'Proyek pertama')
        second = self.ref(2, 'Proyek kedua')
        history = [
            {'id': 'E001', 'project': 'Proyek pertama', 'source_refs': [first], 'supporting_facts': []},
            {'id': 'E002', 'project': 'Proyek kedua', 'source_refs': [second], 'supporting_facts': []},
        ]
        cross = evaluate_cv_supporting_documents({}, history, [], [])
        result = evaluate_anomalies(
            {'identity': {}, 'employment_history': history, 'education_records': [],
             'certificates': [], 'chronology': {}}, {}, cross)
        anomaly = next(item for item in result['items'] if item['experience_ids'] == ['E001'])
        self.assertEqual(anomaly['source_refs'], [first])

    def test_missing_certificate_scheme_is_insufficient_evidence_not_conflict(self):
        ref = self.ref(3, 'CERT-001 Level 7')
        cross = evaluate_cv_supporting_documents(
            {'certificate_summary': 'Ahli Teknik Jembatan Level 7'}, [], [],
            [{'scheme': '', 'level': '7', 'source_refs': [ref]}])
        item = cross['details'][0]['fields'][0]
        self.assertEqual(item['status'], 'missing_support')
        self.assertEqual(cross['conflicts'], [])
        self.assertEqual(cross['missing'], [item])

    def test_partial_certificate_inventory_is_unresolved_not_conflict(self):
        known = self.ref(3, 'Sertifikat Manajemen Konstruksi')
        unknown = self.ref(4, 'Nomor sertifikat tidak terbaca')
        cross = evaluate_cv_supporting_documents(
            {'certificate_summary': 'Ahli Teknik Jembatan Level 7'}, [], [],
            [{'scheme': 'Manajemen Konstruksi', 'level': '7', 'source_refs': [known]},
             {'scheme': '', 'level': '', 'source_refs': [unknown]}])
        item = cross['details'][0]['fields'][0]
        self.assertEqual(item['status'], 'incomplete_mapping')
        self.assertEqual(cross['conflicts'], [])
        self.assertEqual(cross['missing'], [item])

    def test_complete_certificate_inventory_can_confirm_conflict(self):
        ref = self.ref(3, 'Sertifikat Manajemen Konstruksi Level 7')
        cross = evaluate_cv_supporting_documents(
            {'certificate_summary': 'Ahli Teknik Jembatan Level 7'}, [], [],
            [{'scheme': 'Manajemen Konstruksi', 'level': '7', 'source_refs': [ref]}])
        self.assertEqual(cross['details'][0]['fields'][0]['status'], 'conflict')
        self.assertEqual(len(cross['conflicts']), 1)

    def test_empty_history_does_not_turn_claimed_duration_into_zero_mismatch(self):
        cross = evaluate_cv_supporting_documents({}, [], [], [])
        result = evaluate_anomalies(
            {'identity': {'claimed_months': 120}, 'employment_history': [],
             'education_records': [], 'certificates': [],
             'chronology': {'calendar_months_unique': 0, 'uncertain_entries': []}}, {}, cross)
        self.assertFalse(any(item['code'] == 'claimed_duration_mismatch' for item in result['items']))

    def test_partial_history_signal_preserves_records_and_requires_clarification(self):
        rid = self.prepare_run()
        person = self.person()
        person['history_extraction'] = {
            'status': 'partially_extracted', 'identified_count': 2,
            'analysis': 'CV menunjukkan pengalaman tambahan yang belum dapat dipetakan tanpa menebak rinciannya.',
            'source_refs': [self.ref(1, 'Proyek Jembatan Sungai')],
        }
        review.dispatch({'action': 'save_person', 'run_id': rid, 'payload': person})
        stored = storage.load(storage.run_dir(rid) / 'people' / 'P1.json')
        self.assertEqual(len(stored['employment_history']), 1)
        self.assertEqual(stored['history_extraction']['status'], 'partially_extracted')
        anomaly = next(item for item in stored['anomalies'] if item['code'] == 'history_extraction_incomplete')
        self.assertTrue(anomaly['material'])

    def test_new_payload_rejects_no_experience_status_with_positive_claim(self):
        rid = self.prepare_run()
        person = self.person()
        person['identity']['claimed_months'] = 120
        person['employment_history'] = []
        person['employer_checks'] = []
        person['semantic_assessments'] = [
            {**row, 'status': 'perlu_klarifikasi', 'experience_ids': [],
             'clarification': 'Tidak ada riwayat pengalaman untuk dinilai terhadap requirement.'}
            for row in person['semantic_assessments']
        ]
        person['history_extraction'] = {
            'status': 'no_experience_declared', 'identified_count': 0,
            'analysis': 'Payload menyatakan tidak ada pengalaman meskipun durasi positif diklaim.',
            'source_refs': [self.ref(1, 'Personel Final')],
        }
        with self.assertRaisesRegex(ValueError, 'no_experience_declared.*claimed_months'):
            review.dispatch({'action': 'save_person', 'run_id': rid, 'payload': person})

    def test_grounded_zero_documented_experience_is_distinct_from_missing_extraction(self):
        rid = self.prepare_run()
        person = self.person()
        person['identity']['claimed_months'] = 0
        person['employment_history'] = []
        person['employer_checks'] = []
        person['semantic_assessments'] = [
            {**row, 'status': 'perlu_klarifikasi', 'experience_ids': [],
             'clarification': 'Tidak ada riwayat pengalaman untuk dinilai terhadap requirement.'}
            for row in person['semantic_assessments']
        ]
        person['history_extraction'] = {
            'status': 'no_experience_declared', 'identified_count': 0,
            'analysis': 'Halaman CV yang ditinjau tidak menyatakan riwayat pengalaman kerja.',
            'source_refs': [self.ref(1, 'Personel Final')],
        }
        review.dispatch({'action': 'save_person', 'run_id': rid, 'payload': person})
        snapshot = review.snapshot(storage.run_dir(rid))
        duration = snapshot['people'][0]['chronology']
        self.assertEqual(duration['calendar_months_unique'], 0)
        self.assertEqual(snapshot['people'][0]['history_extraction']['status'], 'no_experience_declared')
        exported = review.dispatch({'action': 'export', 'run_id': rid})
        self.assertTrue(exported['success'], exported)

    def test_legacy_contradictory_absence_is_qualified_without_rewriting_checkpoint(self):
        rid = self.prepare_run()
        review.dispatch({'action': 'save_person', 'run_id': rid, 'payload': self.person()})
        path = storage.run_dir(rid) / 'people' / 'P1.json'
        legacy = storage.load(path)
        legacy['identity']['claimed_months'] = 120
        legacy['history_extraction'] = {
            'status': 'no_experience_declared', 'identified_count': 0,
            'analysis': 'Legacy contradictory absence state retained only for compatibility testing.',
            'source_refs': [],
        }
        storage.save(path, legacy)
        snapshot = review.snapshot(storage.run_dir(rid))
        self.assertEqual(snapshot['people'][0]['history_extraction']['status'], 'unknown')
        self.assertEqual(len(snapshot['people'][0]['employment_history']), 1)
        self.assertEqual(storage.load(path)['history_extraction']['status'], 'no_experience_declared')
        first_export = review.dispatch({'action': 'export', 'run_id': rid})
        second_export = review.dispatch({'action': 'export', 'run_id': rid})
        self.assertTrue(first_export['success'], first_export)
        self.assertTrue(second_export['success'], second_export)
        self.assertEqual(len(first_export['artifacts']), 2)
        self.assertEqual(len(second_export['artifacts']), 2)
        self.assertNotEqual(first_export['artifacts'][0]['path'], second_export['artifacts'][0]['path'])
        self.assertEqual(storage.load(path)['history_extraction']['status'], 'no_experience_declared')

    def test_unknown_extraction_with_year_dates_is_idempotent_and_exports(self):
        rid = self.prepare_run()
        person = self.person()
        person['employment_history'][0].update(start_date='2020', end_date='2021')
        person['history_extraction'] = {
            'status': 'unknown', 'identified_count': None,
            'analysis': 'Kelengkapan pemetaan riwayat belum dapat dipastikan dari dokumen sumber.',
            'source_refs': [],
        }
        review.dispatch({'action': 'save_person', 'run_id': rid, 'payload': person})
        root = storage.run_dir(rid)
        stored = storage.load(root / 'people' / 'P1.json')
        self.assertEqual(stored['history_extraction']['status'], 'unknown')
        status = review.dispatch({'action': 'status', 'run_id': rid})
        self.assertEqual(status['saved_person_ids'], ['P1'])
        first = review.snapshot(root)['people'][0]
        second = review.snapshot(root)['people'][0]
        self.assertEqual(first, second)
        self.assertEqual(first['history_extraction']['status'], 'unknown')
        self.assertTrue(first['chronology']['uncertain_entries'])
        exported = review.dispatch({'action': 'export', 'run_id': rid})
        self.assertTrue(exported['success'], exported)
        self.assertEqual(len(exported['artifacts']), 2)

    def test_partial_history_does_not_create_claimed_duration_mismatch(self):
        history = [{'id': 'E001', 'project': 'Proyek terpetakan', 'employer': 'PT Contoh',
                    'start_date': '2020-01', 'end_date': '2020-12', 'source_refs': [],
                    'supporting_refs': [], 'supporting_facts': []}]
        person = {'identity': {'claimed_months': 120}, 'employment_history': history,
                  'history_extraction': {'status': 'partially_extracted', 'identified_count': 5,
                                         'analysis': 'Empat pengalaman lain belum dipetakan.', 'source_refs': []},
                  'education_records': [], 'certificates': [],
                  'chronology': {'calendar_months_unique': 12, 'uncertain_entries': []}}
        result = evaluate_anomalies(person, {}, evaluate_cv_supporting_documents({}, history, [], []))
        codes = {item['code'] for item in result['items']}
        self.assertIn('history_extraction_incomplete', codes)
        self.assertNotIn('claimed_duration_mismatch', codes)

    def test_complete_history_still_reports_genuine_duration_mismatch(self):
        history = [{'id': 'E001', 'project': 'Proyek lengkap', 'employer': 'PT Contoh',
                    'start_date': '2020-01', 'end_date': '2020-12', 'source_refs': [],
                    'supporting_refs': [], 'supporting_facts': []}]
        person = {'identity': {'claimed_months': 120}, 'employment_history': history,
                  'history_extraction': {'status': 'complete', 'identified_count': 1,
                                         'analysis': 'Seluruh riwayat dipetakan.', 'source_refs': []},
                  'education_records': [], 'certificates': [],
                  'chronology': {'calendar_months_unique': 12, 'uncertain_entries': []}}
        result = evaluate_anomalies(person, {}, evaluate_cv_supporting_documents({}, history, [], []))
        self.assertIn('claimed_duration_mismatch', {item['code'] for item in result['items']})

    def test_legacy_incomplete_dates_checkpoint_is_read_without_rewriting_file(self):
        rid = self.prepare_run()
        person = self.person()
        review.dispatch({'action': 'save_person', 'run_id': rid, 'payload': person})
        path = storage.run_dir(rid) / 'people' / 'P1.json'
        legacy = storage.load(path)
        legacy['history_extraction']['status'] = 'incomplete_dates'
        storage.save(path, legacy)
        reviewed = review.snapshot(storage.run_dir(rid))['people'][0]
        self.assertEqual(reviewed['history_extraction']['status'], 'unknown')
        self.assertEqual(storage.load(path)['history_extraction']['status'], 'incomplete_dates')

    def test_export_contains_cross_check_anomaly_and_final_conclusion(self):
        rid = self.prepare_run()
        review.dispatch({'action': 'save_person', 'run_id': rid, 'payload': self.person()})
        exported = review.dispatch({'action': 'export', 'run_id': rid})
        self.assertTrue(exported['success'], exported)
        self.assertEqual(exported['conclusions'][0]['label'], 'Memenuhi')

        book = load_workbook(exported['artifacts'][0]['path'], read_only=True, data_only=True)
        try:
            self.assertIn('Pencocokan Dokumen', book.sheetnames)
            self.assertIn('Konfirmasi Data', book.sheetnames)
            review_sheet = book['Ringkasan Pemeriksaan']
            self.assertEqual(review_sheet['D2'].value, 'Memenuhi')
            audit = book['Hasil Pemeriksaan']
            labels = {audit.cell(row, 4).value for row in range(2, audit.max_row + 1)}
            self.assertIn('Kesesuaian CV dengan dokumen pendukung', labels)
            self.assertIn('Informasi yang perlu diperiksa kembali', labels)
            self.assertIn('Kesimpulan pemenuhan KAK', labels)
        finally:
            book.close()


if __name__ == '__main__':
    unittest.main()
