"""Structured KAK-matching regressions for Scanner Data Diri 0.6.0."""
import copy
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from scanner import review, storage
from scanner.kak_matching import applicable, certificate_details, duration_details, education_details


class KakMatchingV060Test(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.root = Path(temp.name)
        env = patch.dict(os.environ, HERMES_SCANNER_OUTPUT_DIR=str(self.root / 'output'))
        env.start()
        self.addCleanup(env.stop)

    @staticmethod
    def ref(page, quote):
        return {'document_id': 'D001', 'page': page, 'quote': quote}

    def prepare_run(self, mutate_plan=None, minimum_months=12):
        source = self.root / 'structured-kak.pdf'
        source.write_bytes(b'Synthetic PDF; OCR mocked')
        started = review.dispatch({'action': 'start', 'payload': {
            'project': 'Paket KAK Sintetis 0.6',
            'assessment_date': '2026-09-30',
            'allow_web': False,
            'documents': [{'kind': 'cv', 'path': str(source)}],
        }})
        pages = [
            'Personel KAK\nS1 Teknik Sipil\nSite Engineer\nProyek Jembatan Sungai\n'
            'Konsultan: PT Konsultan A | Pemberi Kerja: Dinas PU\n'
            'Kontraktor: PT Kontraktor A\nPengawasan struktur jembatan\n2020-01 sampai 2021-12',
            'IJAZAH\nPersonel KAK\nIjazah S1 Teknik Sipil\nSarjana Teknik\nUniversitas Contoh',
            'CERT-001 Personel KAK Ahli Teknik Jembatan BNSP Level 7',
            'KAK POSISI ENGINEER\n'
            'Posisi Engineer membutuhkan pengalaman sebagai Site Engineer\n'
            'Pengalaman proyek jembatan\n'
            'Pengalaman pada konsultan supervisi\n'
            f'Pengalaman minimal {minimum_months} bulan\n'
            'Uraian tugas pengawasan struktur\n'
            'Pendidikan minimal S1 Teknik Sipil\n'
            'SKK Ahli Teknik Jembatan jenjang 7 masih berlaku\n'
            'Kompetensi teknis pengawasan struktur',
            'SURAT PENGALAMAN\nProyek Jembatan Sungai\nPT Konsultan A\nSite Engineer\n'
            'Pemberi Kerja Dinas PU\nKonsultan PT Konsultan A\nKontraktor PT Kontraktor A\n'
            '2020-01 sampai 2021-12',
        ]

        def extract(path, output, model):
            self.assertTrue(path.samefile(source))
            for index, text in enumerate(pages, 1):
                (output / f'page_{index:04d}.md').write_text(text, encoding='utf-8')
                (output / f'page_{index:04d}.json').write_text('{}', encoding='utf-8')
            return len(pages)

        with patch('scanner.ocr.run_mineru', side_effect=extract):
            result = review.dispatch({'action': 'document', 'run_id': started['run_id'],
                                      'payload': {'document_id': 'D001'}})
        self.assertTrue(result['success'], result)

        plan = {
            'kak': {'title': 'KAK Posisi Engineer', 'version': '1', 'package_id': 'PKT-060',
                    'analysis': 'KAK sintetis memuat requirement atomik untuk pengujian structured matching 0.6.',
                    'receipt_ids': []},
            'requirements': [
                {'id': 'R1', 'role': 'Engineer', 'kind': 'experience',
                 'audit_code': 'position_experience_match',
                 'text': 'Posisi Engineer membutuhkan pengalaman sebagai Site Engineer',
                 'minimum_months': None,
                 'source_refs': [self.ref(4, 'Posisi Engineer membutuhkan pengalaman sebagai Site Engineer')]},
                {'id': 'R2', 'role': 'Engineer', 'kind': 'experience',
                 'audit_code': 'project_kak_match',
                 'text': 'Pengalaman proyek jembatan', 'minimum_months': None,
                 'source_refs': [self.ref(4, 'Pengalaman proyek jembatan')]},
                {'id': 'R3', 'role': 'Engineer', 'kind': 'experience',
                 'audit_code': 'organization_role_match',
                 'text': 'Pengalaman pada konsultan supervisi', 'minimum_months': None,
                 'source_refs': [self.ref(4, 'Pengalaman pada konsultan supervisi')]},
                {'id': 'R5', 'role': 'Engineer', 'kind': 'experience',
                 'audit_code': 'experience_duration_match',
                 'text': f'Pengalaman minimal {minimum_months} bulan', 'minimum_months': minimum_months,
                 'source_refs': [self.ref(4, f'Pengalaman minimal {minimum_months} bulan')]},
                {'id': 'R6', 'role': 'Engineer', 'kind': 'experience',
                 'audit_code': 'responsibility_position_match',
                 'text': 'Uraian tugas pengawasan struktur', 'minimum_months': None,
                 'source_refs': [self.ref(4, 'Uraian tugas pengawasan struktur')]},
                {'id': 'R7', 'role': 'Engineer', 'kind': 'education',
                 'audit_code': 'education_major_match',
                 'text': 'Pendidikan minimal S1 Teknik Sipil', 'minimum_months': None,
                 'parameters': {'minimum_level': 'S1', 'accepted_majors': ['Teknik Sipil']},
                 'source_refs': [self.ref(4, 'Pendidikan minimal S1 Teknik Sipil')]},
                {'id': 'R8', 'role': 'Engineer', 'kind': 'certificate',
                 'audit_code': 'certificate_kak_validity',
                 'text': 'SKK Ahli Teknik Jembatan jenjang 7 masih berlaku', 'minimum_months': None,
                 'parameters': {'schemes': ['Ahli Teknik Jembatan'], 'levels': ['7'],
                                'issuers': [], 'require_current': True},
                 'source_refs': [self.ref(4, 'SKK Ahli Teknik Jembatan jenjang 7 masih berlaku')]},
                {'id': 'R12', 'role': 'Engineer', 'kind': 'other',
                 'audit_code': 'technical_competency_match',
                 'text': 'Kompetensi teknis pengawasan struktur', 'minimum_months': None,
                 'source_refs': [self.ref(4, 'Kompetensi teknis pengawasan struktur')]},
            ],
            'roster': [{'id': 'P1', 'name': 'Personel KAK', 'role': 'Engineer',
                        'cv_refs': [self.ref(1, 'Personel KAK')],
                        'certificate_inventory': [{'id': 'C1', 'source_refs': [
                            self.ref(3, 'CERT-001 Personel KAK Ahli Teknik Jembatan BNSP Level 7')]}]}],
            'coverage': [
                {'document_id': 'D001', 'first_page': 1, 'last_page': 3, 'person_ids': ['P1']},
                {'document_id': 'D001', 'first_page': 4, 'last_page': 4, 'person_ids': [], 'reason': 'reference'},
                {'document_id': 'D001', 'first_page': 5, 'last_page': 5, 'person_ids': ['P1']},
            ],
            'page_classification': [
                {'document_id': 'D001', 'first_page': 1, 'last_page': 1, 'kind': 'cv'},
                {'document_id': 'D001', 'first_page': 2, 'last_page': 3, 'kind': 'attachment'},
                {'document_id': 'D001', 'first_page': 4, 'last_page': 4, 'kind': 'kak'},
                {'document_id': 'D001', 'first_page': 5, 'last_page': 5, 'kind': 'attachment'},
            ],
        }
        if mutate_plan:
            mutate_plan(plan)
        review.dispatch({'action': 'plan', 'run_id': started['run_id'], 'payload': plan})
        return started['run_id']

    def person_payload(self):
        semantic = [
            ('R1', 'Site Engineer'),
            ('R2', 'Proyek Jembatan Sungai'),
            ('R3', 'Konsultan: PT Konsultan A | Pemberi Kerja: Dinas PU'),
            ('R6', 'Pengawasan struktur jembatan'),
            ('R12', 'Pengawasan struktur jembatan'),
        ]
        return {
            'id': 'P1',
            'identity': {'education': 'S1 Teknik Sipil', 'certificate_summary': 'Ahli Teknik Jembatan Level 7',
                         'claimed_months': 24, 'nik': ''},
            'summary': 'Personel sintetis memiliki satu pengalaman, pendidikan, sertifikat, dan requirement KAK terstruktur untuk pengujian matching.',
            'checks': [],
            'education_records': [{
                'level': 'S1', 'degree': 'Sarjana Teknik', 'major': 'Teknik Sipil',
                'institution': 'Universitas Contoh',
                'source_refs': [self.ref(1, 'S1 Teknik Sipil')],
                'supporting_refs': [
                    self.ref(2, 'Ijazah S1 Teknik Sipil'),
                    self.ref(2, 'Sarjana Teknik'),
                    self.ref(2, 'Universitas Contoh'),
                ],
            }],
            'employment_history': [{
                'employer': 'PT Konsultan A', 'role': 'Site Engineer',
                'client': 'Dinas PU', 'consultant': 'PT Konsultan A',
                'contractor': 'PT Kontraktor A', 'represented_organization': '',
                'project': 'Proyek Jembatan Sungai',
                'start_date': '2020-01', 'end_date': '2021-12',
                'relevant': False, 'responsibilities': 'Pengawasan struktur jembatan',
                'source_refs': [
                    self.ref(1, 'Site Engineer'),
                    self.ref(1, 'Proyek Jembatan Sungai'),
                    self.ref(1, 'Konsultan: PT Konsultan A | Pemberi Kerja: Dinas PU'),
                    self.ref(1, 'Pengawasan struktur jembatan'),
                ],
                'supporting_refs': [self.ref(5, 'SURAT PENGALAMAN')],
                'supporting_facts': [{
                    'project': 'Proyek Jembatan Sungai', 'employer': 'PT Konsultan A',
                    'role': 'Site Engineer', 'client': 'Dinas PU', 'consultant': 'PT Konsultan A',
                    'contractor': 'PT Kontraktor A', 'strict_evidence': True,
                    'start_date': '2020-01', 'end_date': '2021-12',
                    'source_refs': [
                        self.ref(5, 'Proyek Jembatan Sungai'),
                        self.ref(5, 'PT Konsultan A'),
                        self.ref(5, 'Site Engineer'),
                        self.ref(5, 'Dinas PU'),
                        self.ref(5, 'PT Kontraktor A'),
                        self.ref(5, '2020-01 sampai 2021-12'),
                    ],
                }],
            }],
            'semantic_assessments': [{
                'requirement_id': requirement_id, 'status': 'memenuhi',
                'finding': f'Bukti kandidat mendukung requirement {requirement_id}.',
                'analysis': 'Klaim pengalaman dibandingkan dengan requirement KAK yang atomik dan bukti halaman kandidat.',
                'experience_ids': ['E001'], 'source_refs': [self.ref(1, quote)],
            } for requirement_id, quote in semantic],
            'employer_checks': [{
                'employer': 'PT Konsultan A',
                'analysis': 'Fixture lokal tidak menguji keberadaan perusahaan melalui web.',
            }],
            'certificates': [{
                'inventory_id': 'C1', 'number': 'CERT-001', 'holder': 'Personel KAK',
                'scheme': 'Ahli Teknik Jembatan', 'issuer': 'BNSP', 'level': '7',
                'issued_on': '2025-01-01', 'expires_on': '2030-12-31',
                'analysis': 'Sertifikat sintetis memuat identitas, skema, penerbit, jenjang, dan masa berlaku untuk matching KAK.',
                'source_refs': [self.ref(3, 'CERT-001 Personel KAK Ahli Teknik Jembatan BNSP Level 7')],
            }],
            'certificate_limitation': '',
            'findings': [],
            'audit_checks': [],
        }

    def test_full_structured_matching_produces_eight_kak_audits(self):
        rid = self.prepare_run()
        result = review.dispatch({'action': 'save_person', 'run_id': rid, 'payload': self.person_payload()})
        statuses = {item['code']: item['status'] for item in result['audit_checks']}
        expected = {
            'position_experience_match', 'project_kak_match', 'organization_role_match',
            'experience_duration_match', 'responsibility_position_match', 'education_major_match',
            'certificate_kak_validity', 'technical_competency_match',
        }
        self.assertTrue(expected <= statuses.keys())
        self.assertTrue(all(statuses[code] == 'memenuhi' for code in expected), statuses)
        self.assertEqual(result['kak_match_detail_count'], 8)

        stored = storage.load(storage.run_dir(rid) / 'people/P1.json')
        self.assertTrue(stored['employment_history'][0]['relevant'])
        self.assertEqual(stored['chronology']['relevant_months'], 24)
        self.assertEqual(len(stored['semantic_assessments']), 5)
        self.assertEqual(len(stored['kak_match_details']), 8)

        snapshot = review.snapshot(storage.run_dir(rid))
        self.assertFalse(any('tidak ada kriteria KAK yang cocok' in item for item in snapshot['limitations']))

    def test_wildcard_role_applies_to_any_role(self):
        requirement = {'role': '*'}
        self.assertTrue(applicable(requirement, 'Engineer'))
        self.assertTrue(applicable(requirement, 'Arsitek'))
        self.assertTrue(applicable({'role': 'Engineer'}, 'Engineer'))
        self.assertFalse(applicable({'role': 'Engineer'}, 'Arsitek'))

    def test_wildcard_duration_requirement_survives_save_snapshot_and_export(self):
        def wildcard_duration(plan):
            next(row for row in plan['requirements'] if row['id'] == 'R5')['role'] = '*'

        rid = self.prepare_run(wildcard_duration)
        saved = review.dispatch({'action': 'save_person', 'run_id': rid, 'payload': self.person_payload()})
        self.assertEqual(
            next(row for row in saved['audit_checks'] if row['code'] == 'experience_duration_match')['status'],
            'memenuhi',
        )
        snapshot = review.snapshot(storage.run_dir(rid))
        self.assertEqual(
            next(row for row in snapshot['people'][0]['kak_match_details'] if row['requirement_id'] == 'R5')['status'],
            'memenuhi',
        )
        exported = review.dispatch({'action': 'export', 'run_id': rid})
        self.assertTrue(exported['success'], exported)
        self.assertEqual(len(exported['artifacts']), 2)

    def test_complete_review_with_zero_relevant_months_exports_failure_with_evidence(self):
        rid = self.prepare_run(minimum_months=60)
        person = self.person_payload()
        position = next(row for row in person['semantic_assessments'] if row['requirement_id'] == 'R1')
        position.update(status='tidak_memenuhi', finding='Pengalaman tidak relevan dengan posisi KAK.')
        person['history_extraction'] = {
            'status': 'complete', 'identified_count': 1,
            'analysis': 'Seluruh pengalaman pada CV telah dipetakan dan diperiksa terhadap posisi KAK.',
            'source_refs': [self.ref(1, 'Proyek Jembatan Sungai')],
        }
        saved = review.dispatch({'action': 'save_person', 'run_id': rid, 'payload': person})
        duration = next(row for row in saved['audit_checks'] if row['code'] == 'experience_duration_match')
        self.assertEqual(duration['status'], 'tidak_memenuhi')
        stored = storage.load(storage.run_dir(rid) / 'people/P1.json')
        stored_duration = next(row for row in stored['audit_checks'] if row['code'] == 'experience_duration_match')
        self.assertTrue(stored_duration['source_refs'])
        self.assertEqual(saved['final_conclusion']['status'], 'tidak_memenuhi')
        self.assertEqual(
            review.snapshot(storage.run_dir(rid))['people'][0]['final_conclusion']['status'],
            'tidak_memenuhi',
        )
        exported = review.dispatch({'action': 'export', 'run_id': rid})
        self.assertTrue(exported['success'], exported)
        self.assertEqual(len(exported['artifacts']), 2)

    def test_export_contains_matching_and_education_sheets(self):
        rid = self.prepare_run()
        review.dispatch({'action': 'save_person', 'run_id': rid, 'payload': self.person_payload()})
        result = review.dispatch({'action': 'export', 'run_id': rid})
        self.assertTrue(result['success'], result)
        book = load_workbook(result['artifacts'][0]['path'], read_only=True, data_only=True)
        try:
            self.assertIn('Pemeriksaan KAK Terperinci', book.sheetnames)
            self.assertIn('Pendidikan', book.sheetnames)
            detail = book['Pemeriksaan KAK Terperinci']
            self.assertEqual(detail.max_row, 9)
            labels = {detail.cell(row, 4).value for row in range(2, detail.max_row + 1)}
            self.assertIn('Kesesuaian kompetensi teknis', labels)
            self.assertNotIn('technical_competency_match', labels)
            education = book['Pendidikan']
            self.assertEqual(education['D2'].value, 'S1')
            history = book['Riwayat Pekerjaan']
            self.assertEqual(history['G2'].value, 'Dinas PU')
        finally:
            book.close()

    def test_missing_semantic_requirement_is_rejected(self):
        rid = self.prepare_run()
        person = self.person_payload()
        person['semantic_assessments'].pop()
        with self.assertRaisesRegex(ValueError, 'tepat semua requirement semantic'):
            review.dispatch({'action': 'save_person', 'run_id': rid, 'payload': person})

    def test_semantic_assessment_cannot_reference_unknown_experience(self):
        rid = self.prepare_run()
        person = self.person_payload()
        person['semantic_assessments'][0]['experience_ids'] = ['E999']
        with self.assertRaisesRegex(ValueError, 'pengalaman tidak dikenal'):
            review.dispatch({'action': 'save_person', 'run_id': rid, 'payload': person})

    def test_invented_structured_parameter_is_rejected_at_plan(self):
        def mutate(plan):
            requirement = next(row for row in plan['requirements'] if row['id'] == 'R7')
            requirement['parameters']['accepted_majors'] = ['Teknik Mesin']
        with self.assertRaisesRegex(ValueError, 'tidak ditemukan pada kutipan'):
            self.prepare_run(mutate)

    def test_duration_below_minimum_is_not_memenuhi_when_period_is_complete(self):
        requirement = {'id': 'D', 'role': 'Engineer', 'audit_code': 'experience_duration_match',
                       'text': 'Minimal 36 bulan', 'minimum_months': 36, 'source_refs': []}
        history = [{'id': 'E001', 'relevant': True, 'source_refs': [], 'supporting_refs': []}]
        chronology = {'relevant_months': 24, 'supported_relevant_months': 24, 'uncertain_entries': []}
        detail = duration_details([requirement], 'Engineer', history, chronology, {
            'history_extraction': {'status': 'complete'}, 'relevance_assessed': True,
        })[0]
        self.assertEqual(detail['status'], 'tidak_memenuhi')

    def test_duration_below_minimum_is_unresolved_when_history_is_partial_or_unknown(self):
        requirement = {'id': 'D', 'role': 'Engineer', 'audit_code': 'experience_duration_match',
                       'text': 'Minimal 60 bulan', 'minimum_months': 60, 'source_refs': []}
        history = [{'id': 'E001', 'relevant': True, 'source_refs': [], 'supporting_refs': []}]
        chronology = {'relevant_months': 12, 'supported_relevant_months': 12, 'uncertain_entries': []}
        for status in ('partially_extracted', 'unknown'):
            with self.subTest(status=status):
                detail = duration_details([requirement], 'Engineer', history, chronology, {
                    'history_extraction': {'status': status}, 'relevance_assessed': True,
                })[0]
                self.assertEqual(detail['status'], 'perlu_klarifikasi')

    def test_supported_subset_can_establish_minimum_despite_partial_history(self):
        requirement = {'id': 'D', 'role': 'Engineer', 'audit_code': 'experience_duration_match',
                       'text': 'Minimal 12 bulan', 'minimum_months': 12, 'source_refs': []}
        history = [{'id': 'E001', 'relevant': True, 'source_refs': [], 'supporting_refs': []}]
        chronology = {'relevant_months': 12, 'supported_relevant_months': 12, 'uncertain_entries': []}
        detail = duration_details([requirement], 'Engineer', history, chronology, {
            'history_extraction': {'status': 'partially_extracted'}, 'relevance_assessed': True,
        })[0]
        self.assertEqual(detail['status'], 'memenuhi')

    def test_missing_evaluation_context_is_conservative(self):
        requirement = {'id': 'D', 'role': 'Engineer', 'audit_code': 'experience_duration_match',
                       'text': 'Minimal 60 bulan', 'minimum_months': 60, 'source_refs': []}
        detail = duration_details([requirement], 'Engineer', [], {
            'relevant_months': 0, 'supported_relevant_months': 0, 'uncertain_entries': [],
        })[0]
        self.assertEqual(detail['status'], 'perlu_klarifikasi')

    def test_zero_month_requirement_does_not_claim_verified_experience(self):
        requirement = {'id': 'D', 'role': 'Engineer', 'audit_code': 'experience_duration_match',
                       'text': 'Minimal 0 bulan', 'minimum_months': 0, 'source_refs': []}
        detail = duration_details([requirement], 'Engineer', [], {
            'relevant_months': 0, 'supported_relevant_months': 0, 'uncertain_entries': [],
        })[0]
        self.assertEqual(detail['status'], 'memenuhi')
        self.assertIn('bukan bukti', detail['finding'])
        self.assertFalse(detail['candidate_evidence_required'])

    def test_zero_month_requirement_survives_validation_and_export_without_claiming_verification(self):
        def duration_only(plan):
            plan['requirements'] = [row for row in plan['requirements'] if row['id'] == 'R5']
            plan['roster'][0]['certificate_inventory'] = []

        rid = self.prepare_run(duration_only, minimum_months=0)
        person = self.person_payload()
        person.update(
            identity={'education': '', 'certificate_summary': '', 'claimed_months': 0, 'nik': ''},
            employment_history=[], education_records=[], semantic_assessments=[], employer_checks=[], certificates=[],
            certificate_limitation='Tidak ada sertifikat yang dipetakan karena requirement ini hanya menetapkan ambang durasi nol.',
        )
        saved = review.dispatch({'action': 'save_person', 'run_id': rid, 'payload': person})
        stored = storage.load(storage.run_dir(rid) / 'people/P1.json')
        detail = next(row for row in stored['kak_match_details'] if row['requirement_id'] == 'R5')
        self.assertEqual(detail['status'], 'memenuhi')
        self.assertIn('bukan bukti', detail['finding'])
        self.assertFalse(detail['candidate_evidence_required'])
        exported = review.dispatch({'action': 'export', 'run_id': rid})
        self.assertTrue(exported['success'], (saved, exported))
        self.assertEqual(len(exported['artifacts']), 2)

    def test_partial_history_below_kak_minimum_stays_clarification_through_export(self):
        rid = self.prepare_run(minimum_months=60)
        person = self.person_payload()
        person['identity']['claimed_months'] = 120
        person['employment_history'][0].update(start_date='2021-01', end_date='2021-12', supporting_facts=[])
        person['history_extraction'] = {
            'status': 'partially_extracted', 'identified_count': 5,
            'analysis': 'CV menunjukkan empat pengalaman tambahan yang belum berhasil dipetakan.',
            'source_refs': [self.ref(1, 'Proyek Jembatan Sungai')],
        }
        saved = review.dispatch({'action': 'save_person', 'run_id': rid, 'payload': person})
        statuses = {item['code']: item['status'] for item in saved['audit_checks']}
        self.assertEqual(statuses['experience_duration_match'], 'perlu_klarifikasi')
        self.assertEqual(saved['final_conclusion']['status'], 'perlu_klarifikasi')
        snap = review.snapshot(storage.run_dir(rid))
        self.assertEqual(snap['people'][0]['final_conclusion']['status'], 'perlu_klarifikasi')
        exported = review.dispatch({'action': 'export', 'run_id': rid})
        self.assertTrue(exported['success'], exported)
        self.assertEqual(len(exported['artifacts']), 2)
        workbook = load_workbook(exported['artifacts'][0]['path'], read_only=True, data_only=True)
        try:
            sheet = workbook['Ringkasan Pemeriksaan']
            headers = [cell.value for cell in sheet[1]]
            conclusion = headers.index('Kesimpulan KAK') + 1
            self.assertEqual(sheet.cell(2, conclusion).value, 'Perlu Klarifikasi')
        finally:
            workbook.close()

    def test_education_level_failure_and_major_ambiguity_are_distinct(self):
        req = {'id': 'E', 'role': 'Engineer', 'audit_code': 'education_major_match',
               'text': 'S1 Teknik Sipil', 'source_refs': [],
               'parameters': {'minimum_level': 'S1', 'accepted_majors': ['Teknik Sipil']}}
        low = education_details([req], 'Engineer', [{
            'level': 'D3', 'major': 'Teknik Sipil', 'source_refs': [], 'supporting_refs': []
        }])[0]
        ambiguous = education_details([req], 'Engineer', [{
            'level': 'S1', 'major': 'Teknik Struktur', 'source_refs': [], 'supporting_refs': []
        }])[0]
        self.assertEqual(low['status'], 'tidak_memenuhi')
        self.assertEqual(ambiguous['status'], 'perlu_klarifikasi')

    def test_expired_matching_certificate_is_not_memenuhi(self):
        req = {'id': 'C', 'role': 'Engineer', 'audit_code': 'certificate_kak_validity',
               'text': 'SKK Level 7 aktif', 'source_refs': [],
               'parameters': {'schemes': ['Ahli Teknik Jembatan'], 'levels': ['7'],
                              'issuers': [], 'require_current': True}}
        cert = {'scheme': 'Ahli Teknik Jembatan', 'level': '7', 'issuer': 'BNSP',
                'validity': 'kedaluwarsa', 'source_refs': []}
        detail = certificate_details([req], 'Engineer', [cert])[0]
        self.assertEqual(detail['status'], 'tidak_memenuhi')

    def test_payload_cannot_override_structured_server_audit(self):
        rid = self.prepare_run()
        person = self.person_payload()
        person['audit_checks'] = [{
            'code': 'education_major_match', 'status': 'memenuhi',
            'finding': 'Dipaksakan payload.', 'analysis': 'Payload mencoba menimpa hasil matching pendidikan server.',
            'source_refs': [self.ref(1, 'S1 Teknik Sipil')], 'kak_refs': [],
        }]
        with self.assertRaisesRegex(ValueError, 'dikelola server'):
            review.dispatch({'action': 'save_person', 'run_id': rid, 'payload': person})


if __name__ == '__main__':
    unittest.main()
