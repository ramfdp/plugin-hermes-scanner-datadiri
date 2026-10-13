"""One bundled PDF, no KAK: real exports with synthetic OCR/web boundaries.

Never use applicant documents or identifiers as repository fixtures.
"""
import asyncio
import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook
from reportlab.pdfgen import canvas
from scanner import __version__, storage, tools, workflow


def source(page, quote):
    return {'document_id': 'D001', 'page': page, 'quote': quote}


class SinglePDFTest(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.root = Path(temp.name)
        env = patch.dict(os.environ, HERMES_SCANNER_OUTPUT_DIR=str(self.root / 'output'))
        env.start()
        self.addCleanup(env.stop)

    def prepare(self, count=9, allow_web=True, employment_evidence=False):
        # Six synthetic pages per person: 54 pages in the principal regression.
        pages, roster, people, coverage, summaries = [], [], [], [], []
        for i in range(count):
            first = 6 * i + 1
            name, pid, number = f'Personel Sintetis {i + 1}', f'P{i + 1}', f'TEST-{i + 1:03d}'
            cv = f'CURRICULUM VITAE\n{name}\nS1 Bidang Contoh\nKlaim pengalaman 24 bulan'
            certificate = f'{number} {name} Skema Contoh Penerbit Contoh'
            if employment_evidence:
                pages.extend([cv, 'Project Alpha PT Example Site Engineer 2020-01 sampai 2021-12',
                              'KAK Engineer pengalaman sebagai Site Engineer minimal 12 bulan',
                              f'Ijazah S1 Bidang Contoh {name}', certificate,
                              'SURAT PENGALAMAN Project Alpha 2021-01 sampai 2021-12'])
            else:
                pages.extend([cv, f'Lanjutan CV {name}', f'Identitas {name} [ID_REDACTED]',
                              f'Ijazah S1 Bidang Contoh {name}', certificate, f'Lampiran kompetensi {name}'])
            refs = [source(first + 4, certificate)]
            roster.append({'id': pid, 'name': name, 'role': 'Engineer', 'cv_refs': [source(first, name)],
                           'certificate_inventory': [{'id': f'C{i + 1}', 'source_refs': refs}]})
            if employment_evidence:
                coverage.extend([
                    {'document_id': 'D001', 'first_page': first, 'last_page': first + 1, 'person_ids': [pid]},
                    {'document_id': 'D001', 'first_page': first + 2, 'last_page': first + 2,
                     'person_ids': [], 'reason': 'reference'},
                    {'document_id': 'D001', 'first_page': first + 3, 'last_page': first + 5, 'person_ids': [pid]},
                ])
            else:
                coverage.append({'document_id': 'D001', 'first_page': first, 'last_page': first + 5, 'person_ids': [pid]})
            identity = {'education': 'S1 Bidang Contoh', 'certificate_summary': 'Skema Contoh',
                        'claimed_months': 24 if employment_evidence else 12, 'nik': ''}
            summaries.append({'id': pid, 'identity': identity.copy(), 'source_refs': [source(first, cv)]})
            person = {'id': pid, 'identity': identity.copy(),
                'summary': 'CV dan lampiran berada pada satu PDF. Klaim pengalaman belum didukung kronologi terperinci; acuan KAK tidak tersedia.',
                'checks': [], 'employment_history': [], 'employer_checks': [],
                'certificates': [{'inventory_id': f'C{i + 1}', 'number': number, 'holder': name,
                    'scheme': 'Skema Contoh', 'issuer': 'Penerbit Contoh', 'issued_on': None, 'expires_on': None,
                    'analysis': 'Nomor dan identitas terbaca pada lampiran; pemeriksaan portal belum memberikan bukti yang cukup untuk menyatakan keabsahan.',
                    'source_refs': refs}], 'certificate_limitation': '',
                'findings': ['Tidak ada acuan KAK; kesesuaian persyaratan belum dapat dinilai.']}
            if employment_evidence:
                person.update(
                    summary='CV, lampiran pengalaman, sertifikat, dan KAK berada pada satu PDF dengan klasifikasi halaman terpisah.',
                    employment_history=[{
                        'project': 'Project Alpha', 'employer': 'PT Example', 'role': 'Site Engineer',
                        'start_date': '2020-01', 'end_date': '2021-12',
                        'responsibilities': 'Pelaksanaan tugas Site Engineer.',
                        'source_refs': [source(first + 1, 'Project Alpha PT Example Site Engineer 2020-01 sampai 2021-12')],
                        'supporting_refs': [source(first + 5, 'SURAT PENGALAMAN Project Alpha 2021-01 sampai 2021-12')],
                        'supporting_facts': [{
                            'project': 'Project Alpha', 'start_date': '2021-01', 'end_date': '2021-12',
                            'strict_evidence': True,
                            'source_refs': [source(first + 5, 'SURAT PENGALAMAN Project Alpha 2021-01 sampai 2021-12')],
                        }],
                    }],
                    history_extraction={
                        'status': 'complete', 'identified_count': 1,
                        'analysis': 'Satu pengalaman pada CV telah dipetakan seluruhnya dari halaman yang ditinjau.',
                        'source_refs': [source(first + 1, 'Project Alpha')],
                    },
                    semantic_assessments=[{
                        'requirement_id': 'RPOS', 'status': 'memenuhi',
                        'finding': 'Posisi Site Engineer ditemukan pada pengalaman Project Alpha.',
                        'analysis': 'Posisi pengalaman dibandingkan langsung dengan requirement posisi KAK.',
                        'experience_ids': ['E001'], 'source_refs': [source(first + 1, 'Site Engineer')],
                    }],
                    employer_checks=[{'employer': 'PT Example',
                                      'analysis': 'Fixture sintetis tidak melakukan verifikasi perusahaan melalui web.'}],
                    findings=[],
                )
            people.append(person)
        path = self.root / 'Gabungan CV sintetis.pdf'
        pdf = canvas.Canvas(str(path))
        for text in pages:
            for line, value in enumerate(text.splitlines()):
                pdf.drawString(35, 800 - line * 20, value)
            pdf.showPage()
        pdf.save()
        result = workflow.dispatch({'action': 'start', 'payload': {
            'workflow_version': __version__, 'project': path.stem, 'assessment_date': '2026-09-23',
            'expected_person_count': None, 'allow_web': allow_web,
            'documents': [{'kind': 'cv', 'path': str(path)}]}})
        rid = result['run_id']
        self.assertEqual(len(result['documents']), 1)

        def extract(input_path, output, model_dir):
            # Windows TEMP may use an 8.3 alias; compare the actual file identity.
            self.assertTrue(input_path.samefile(path))
            for n, text in enumerate(pages, 1):
                (output / f'page_{n:04d}.md').write_text(text, encoding='utf-8')
                (output / f'page_{n:04d}.json').write_text('{}', encoding='utf-8')
            return len(pages)

        with patch('scanner.ocr.run_mineru', side_effect=extract) as ocr:
            result = workflow.dispatch({'action': 'document', 'run_id': rid, 'payload': {'document_id': 'D001'}})
            self.assertTrue(result['success'], result)
            cached = workflow.dispatch({'action': 'document', 'run_id': rid, 'payload': {'document_id': 'D001'}})
            self.assertTrue(cached['cached']); self.assertEqual(ocr.call_count, 1)
        payload, read_pages = {'document_id': 'D001', 'page': 1}, []
        while payload is not None:
            item = workflow.dispatch({'action': 'read_document', 'run_id': rid, 'payload': payload})
            read_pages.append(item['page'])
            payload = item['next']
        self.assertEqual(read_pages, list(range(1, len(pages) + 1)))
        plan = {'kak': {'title': 'KAK Sintetis' if employment_evidence else '', 'version': '',
                       'package_id': '', 'receipt_ids': [],
                       'analysis': ('KAK sintetis memuat persyaratan posisi dan durasi untuk pengujian.' if employment_evidence else
                                    'Tidak ada dokumen KAK. Review CV dan sertifikat dilanjutkan dengan keterbatasan penilaian KAK.')},
                'requirements': ([
                    {'id': 'RPOS', 'role': '*', 'kind': 'experience', 'audit_code': 'position_experience_match',
                     'text': 'pengalaman sebagai Site Engineer', 'minimum_months': None,
                     'source_refs': [source(3, 'pengalaman sebagai Site Engineer')]},
                    {'id': 'RDUR', 'role': '*', 'kind': 'experience', 'audit_code': 'experience_duration_match',
                     'text': 'minimal 12 bulan', 'minimum_months': 12,
                     'source_refs': [source(3, 'minimal 12 bulan')]},
                ] if employment_evidence else []), 'roster': roster, 'coverage': coverage}
        if employment_evidence:
            plan['page_classification'] = [
                {'document_id': 'D001', 'first_page': 1, 'last_page': 2, 'kind': 'cv'},
                {'document_id': 'D001', 'first_page': 3, 'last_page': 3, 'kind': 'kak'},
                {'document_id': 'D001', 'first_page': 4, 'last_page': 6, 'kind': 'attachment'},
            ]
        result = workflow.dispatch({'action': 'plan', 'run_id': rid, 'payload': plan})
        self.assertEqual(result['kak_status'],
                         'dokumen_diberikan_belum_diautentikasi' if employment_evidence else 'tidak_tersedia')
        summary = workflow.dispatch({'action': 'summary', 'run_id': rid, 'payload': {'people': summaries}})
        self.assertTrue(summary['success'], summary)
        self.assertTrue(Path(summary['artifact']['path']).is_file())
        return rid, people, summary

    def test_one_combined_pdf_binds_attachment_period_and_exports_two_final_artifacts(self):
        rid, people, _ = self.prepare(count=1, allow_web=False, employment_evidence=True)
        verified = workflow.dispatch({'action': 'verify_kak', 'run_id': rid, 'payload': {
            'analysis': 'KAK sintetis telah dipetakan dari halaman KAK pada PDF gabungan.', 'receipt_ids': [],
        }})
        self.assertTrue(verified['success'], verified)
        saved = workflow.dispatch({'action': 'save_person', 'run_id': rid, 'payload': people[0]})
        self.assertTrue(saved['success'], saved)
        person = storage.load(storage.run_dir(rid) / 'people/P1.json')
        self.assertEqual(person['chronology']['calendar_months_unique'], 24)
        self.assertEqual(person['chronology']['supported_relevant_months'], 12)
        status = workflow.dispatch({'action': 'status', 'run_id': rid})
        self.assertEqual(status['next_action'], 'export')
        exported = workflow.dispatch({'action': 'export', 'run_id': rid})
        self.assertTrue(exported['workflow_complete'], exported)
        self.assertEqual(len(exported['artifacts']), 2)
        self.assertEqual(storage.load(storage.run_dir(rid) / 'manifest.json')['documents'][0]['kind'], 'cv')

    def test_54_page_single_pdf_no_kak_and_failed_web_still_exports_nine_people(self):
        rid, people, summary = self.prepare()

        class UnavailableWeb:
            def dispatch_tool(self, name, arguments):
                raise RuntimeError('Synthetic portal unavailable; no actual network request')

        handler = tools.ordered_web_handler(UnavailableWeb())
        receipt = json.loads(asyncio.run(handler({'run_id': rid, 'purpose': 'certificate', 'tool': 'web_extract',
                          'arguments': {'urls': ['https://bnsp.go.id/check-certification']}})))
        self.assertFalse(receipt['success'])
        self.assertTrue(receipt['receipt_id'])
        result = workflow.dispatch({'action': 'verify_kak', 'run_id': rid, 'payload': {
            'analysis': 'Tidak ada dokumen KAK pada paket CV ini; penilaian syarat KAK belum dapat dilakukan.', 'receipt_ids': []}})
        self.assertTrue(result['success'], result)
        for person in people:
            result = workflow.dispatch({'action': 'save_person', 'run_id': rid, 'payload': person})
            self.assertEqual(result['overall'], 'perlu_klarifikasi')
            self.assertEqual(result['certificates'][0]['status'], 'belum_dapat_diverifikasi')
        exported = workflow.dispatch({'action': 'export', 'run_id': rid})
        self.assertTrue(exported['workflow_complete'], exported)
        self.assertFalse(exported['report_complete'])
        self.assertEqual(len(exported['artifacts']), 2)
        self.assertEqual(exported['person_count'], 9)
        book = load_workbook(exported['artifacts'][0]['path'], read_only=True, data_only=True)
        try:
            self.assertEqual(book['Ringkasan Personel'].max_row, 14)
            self.assertEqual(book['Ringkasan Personel']['G6'].value, 'Belum tersedia')
            self.assertEqual(book['Ringkasan Personel']['B14'].value, 'Personel Sintetis 9')
            certificates = book['Verifikasi Sertifikat']
            headers = {cell.value: cell.column for cell in certificates[1]}
            self.assertEqual(certificates.cell(2, headers['Verifikasi identitas']).value, 'Belum dapat diverifikasi')
        finally:
            book.close()
        self.assertTrue(Path(exported['artifacts'][1]['path']).is_file())
        self.assertTrue(Path(summary['artifact']['path']).is_file())
        self.assertNotIn('.json', exported['chat_markdown'])
        self.assertEqual(storage.load(storage.run_dir(rid) / 'manifest.json')['documents'][0]['kind'], 'cv')

    def test_one_pdf_no_kak_also_preserves_backend_web_opt_out(self):
        rid, people, _ = self.prepare(count=1, allow_web=False)
        for person in people:
            workflow.dispatch({'action': 'save_person', 'run_id': rid, 'payload': person})
        result = workflow.dispatch({'action': 'export', 'run_id': rid})
        self.assertTrue(result['workflow_complete'], result)
        self.assertFalse(result['report_complete'])
        self.assertTrue(any('tidak diizinkan' in message for message in result['limitations']))
