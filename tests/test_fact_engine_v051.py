"""Deterministic fact-engine regressions for Scanner Data Diri 0.5.1."""
import copy
import os
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from scanner import review, storage
from scanner.evaluation import (
    detect_duplicates,
    detect_overlaps,
    evaluate_consistency,
    evaluate_period_accuracy,
    partial_date,
)


class FactEngineV051Test(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.root = Path(temp.name)
        env = patch.dict(os.environ, HERMES_SCANNER_OUTPUT_DIR=str(self.root / 'output'))
        env.start()
        self.addCleanup(env.stop)

    @staticmethod
    def ref(page, quote):
        return {'document_id': 'D001', 'page': page, 'quote': quote}

    def prepare_run(self):
        source = self.root / 'facts.pdf'
        source.write_bytes(b'Synthetic PDF; OCR mocked')
        started = review.dispatch({'action': 'start', 'payload': {
            'project': 'Paket Fakta Sintetis',
            'assessment_date': '2026-09-30',
            'allow_web': False,
            'documents': [{'kind': 'cv', 'path': str(source)}],
        }})
        pages = [
            'Personel Fakta\nProject Alpha\nPT Contoh\nEngineer\n2020-01 sampai 2020-12',
            'SURAT REFERENSI A\nProject Alpha\nPT Contoh\nEngineer\n2020-01 sampai 2020-12',
            'Project Beta\nPT Contoh\nEngineer\n2020-07 sampai 2021-06',
            'SURAT REFERENSI B\nProject Beta\nPT Contoh\nEngineer\n2020-07 sampai 2021-06',
        ]

        def extract(path, output, model):
            self.assertTrue(path.samefile(source))
            for index, text in enumerate(pages, 1):
                (output / f'page_{index:04d}.md').write_text(text, encoding='utf-8')
                (output / f'page_{index:04d}.json').write_text('{}', encoding='utf-8')
            return len(pages)

        with patch('scanner.ocr.run_mineru', side_effect=extract):
            result = review.dispatch({'action': 'document', 'run_id': started['run_id'],
                                      'payload': {'document_id': 'D001'}})
        self.assertTrue(result['success'], result)

        plan = {
            'kak': {'title': '', 'version': '', 'package_id': '', 'receipt_ids': [],
                    'analysis': 'KAK tidak tersedia pada fixture fakta; audit kronologi tetap dijalankan secara lokal.'},
            'requirements': [],
            'roster': [{'id': 'P1', 'name': 'Personel Fakta', 'role': 'Engineer',
                        'cv_refs': [self.ref(1, 'Personel Fakta')], 'certificate_inventory': []}],
            'coverage': [
                {'document_id': 'D001', 'first_page': 1, 'last_page': 4, 'person_ids': ['P1']},
            ],
            'page_classification': [
                {'document_id': 'D001', 'first_page': 1, 'last_page': 1, 'kind': 'cv'},
                {'document_id': 'D001', 'first_page': 2, 'last_page': 2, 'kind': 'attachment'},
                {'document_id': 'D001', 'first_page': 3, 'last_page': 3, 'kind': 'cv'},
                {'document_id': 'D001', 'first_page': 4, 'last_page': 4, 'kind': 'attachment'},
            ],
        }
        review.dispatch({'action': 'plan', 'run_id': started['run_id'], 'payload': plan})
        return started['run_id']

    def person_payload(self):
        return {
            'id': 'P1',
            'identity': {'education': 'S1 Teknik Sipil', 'certificate_summary': '',
                         'claimed_months': 18, 'nik': ''},
            'summary': 'Dua pengalaman proyek sintetis memiliki periode dan dokumen pendukung terstruktur untuk pengujian engine fakta.',
            'checks': [],
            'employment_history': [
                {'employer': 'PT Contoh', 'role': 'Engineer', 'project': 'Project Alpha',
                 'start_date': '2020-01', 'end_date': '2020-12', 'relevant': True,
                 'responsibilities': 'Pengawasan Project Alpha.',
                 'source_refs': [self.ref(1, 'Project Alpha')],
                 'supporting_refs': [self.ref(2, 'SURAT REFERENSI A')],
                 'supporting_facts': [{
                     'project': 'Project Alpha', 'employer': 'PT Contoh', 'role': 'Engineer',
                     'start_date': '2020-01', 'end_date': '2020-12',
                     'source_refs': [self.ref(2, 'Project Alpha')],
                 }]},
                {'employer': 'PT Contoh', 'role': 'Engineer', 'project': 'Project Beta',
                 'start_date': '2020-07', 'end_date': '2021-06', 'relevant': True,
                 'responsibilities': 'Pengawasan Project Beta.',
                 'source_refs': [self.ref(3, 'Project Beta')],
                 'supporting_refs': [self.ref(4, 'SURAT REFERENSI B')],
                 'supporting_facts': [{
                     'project': 'Project Beta', 'employer': 'PT Contoh', 'role': 'Engineer',
                     'start_date': '2020-07', 'end_date': '2021-06',
                     'source_refs': [self.ref(4, 'Project Beta')],
                 }]},
            ],
            'employer_checks': [{'employer': 'PT Contoh',
                                 'analysis': 'Fixture lokal tidak melakukan verifikasi web perusahaan.'}],
            'certificates': [],
            'certificate_limitation': 'Sertifikat tidak menjadi bagian fixture engine fakta.',
            'findings': [],
            'audit_checks': [],
        }

    def test_period_parser_preserves_precision(self):
        year = partial_date('2020')
        month = partial_date('2020-02')
        day = partial_date('2020-02-29')
        bad = partial_date('2020-02-31')
        self.assertEqual((year['precision'], year['first'].isoformat(), year['last'].isoformat()),
                         ('year', '2020-01-01', '2020-12-31'))
        self.assertEqual((month['precision'], month['first'].isoformat(), month['last'].isoformat()),
                         ('month', '2020-02-01', '2020-02-29'))
        self.assertEqual(day['precision'], 'day')
        self.assertFalse(bad['valid'])

    def test_period_and_consistency_match_supporting_facts(self):
        history = self.person_payload()['employment_history']
        for index, row in enumerate(history, 1):
            row['id'] = f'E{index:03d}'
        period = evaluate_period_accuracy(history, date(2026, 9, 30))
        consistency = evaluate_consistency(history)
        self.assertEqual(period['check']['status'], 'memenuhi')
        self.assertEqual(consistency['check']['status'], 'memenuhi')
        self.assertEqual(period['issues'], [])
        self.assertEqual(consistency['conflicts'], [])

    def test_conflicting_supporting_fact_requires_clarification(self):
        history = self.person_payload()['employment_history'][:1]
        history[0]['id'] = 'E001'
        history[0]['supporting_facts'][0]['start_date'] = '2021-01'
        history[0]['supporting_facts'][0]['project'] = 'Project Yang Berbeda'
        period = evaluate_period_accuracy(history, date(2026, 9, 30))
        consistency = evaluate_consistency(history)
        self.assertEqual(period['check']['status'], 'perlu_klarifikasi')
        self.assertEqual(consistency['check']['status'], 'perlu_klarifikasi')
        self.assertTrue(any(item['type'] == 'period_conflict' for item in period['issues']))
        self.assertEqual({item['field'] for item in consistency['conflicts']}, {'project', 'start_date'})

    def test_overlap_is_flagged_but_not_declared_failure(self):
        history = self.person_payload()['employment_history']
        for index, row in enumerate(history, 1):
            row['id'] = f'E{index:03d}'
        result = detect_overlaps(history)
        self.assertEqual(result['check']['status'], 'perlu_klarifikasi')
        self.assertEqual(len(result['pairs']), 1)
        self.assertEqual((result['pairs'][0]['left_id'], result['pairs'][0]['right_id']), ('E001', 'E002'))
        self.assertEqual(result['pairs'][0]['overlap_start'], '2020-07-01')
        self.assertEqual(result['pairs'][0]['overlap_end'], '2020-12-31')

    def test_exact_and_possible_duplicates_are_separate(self):
        base = {'id': 'E001', 'project': 'Pembangunan Jembatan Sungai A', 'employer': 'PT Contoh',
                'role': 'Engineer', 'start_date': '2020-01', 'end_date': '2020-12',
                'source_refs': [], 'supporting_refs': [], 'supporting_facts': []}
        exact = {**copy.deepcopy(base), 'id': 'E002'}
        possible = {**copy.deepcopy(base), 'id': 'E003', 'project': 'Pembangunan Jembatan Sungai-A',
                    'end_date': '2020-11'}
        result = detect_duplicates([base, exact, possible])
        kinds = [row['type'] for row in result['pairs']]
        self.assertIn('exact', kinds)
        self.assertIn('possible', kinds)
        self.assertEqual(result['check']['status'], 'perlu_klarifikasi')

    def test_server_computes_four_fact_audits_and_renders_them(self):
        rid = self.prepare_run()
        result = review.dispatch({'action': 'save_person', 'run_id': rid, 'payload': self.person_payload()})
        statuses = {row['code']: row['status'] for row in result['audit_checks']}
        self.assertEqual(statuses['project_period_accuracy'], 'memenuhi')
        self.assertEqual(statuses['identity_consistency'], 'memenuhi')
        self.assertEqual(statuses['project_overlap'], 'perlu_klarifikasi')
        self.assertEqual(statuses['project_duplicate'], 'memenuhi')

        stored = storage.load(storage.run_dir(rid) / 'people/P1.json')
        self.assertEqual(stored['fact_analysis']['overlaps']['pairs'][0]['left_id'], 'E001')
        self.assertEqual(stored['employment_history'][1]['id'], 'E002')

        exported = review.dispatch({'action': 'export', 'run_id': rid})
        self.assertTrue(exported['success'], exported)
        book = load_workbook(exported['artifacts'][0]['path'], read_only=True, data_only=True)
        try:
            self.assertIn('Hasil Pemeriksaan', book.sheetnames)
            self.assertIn('Periode dan Pencatatan Ganda', book.sheetnames)
            audit = book['Hasil Pemeriksaan']
            values = [audit.cell(row, 4).value for row in range(2, audit.max_row + 1)]
            self.assertIn('Pengalaman dengan periode bersamaan', values)
            self.assertNotIn('project_overlap', values)
            pairs = book['Periode dan Pencatatan Ganda']
            self.assertEqual(pairs['D2'].value, 'Periode bersamaan')
            self.assertIn('Project Alpha', pairs['E2'].value)
        finally:
            book.close()

    def test_payload_cannot_spoof_server_computed_fact_audit(self):
        rid = self.prepare_run()
        person = self.person_payload()
        person['audit_checks'] = [{
            'code': 'project_overlap', 'status': 'memenuhi',
            'finding': 'Tidak ada overlap.', 'analysis': 'Payload mencoba menimpa hasil backend.',
            'source_refs': [self.ref(1, 'Project Alpha')], 'kak_refs': [],
        }]
        with self.assertRaisesRegex(ValueError, 'server-side'):
            review.dispatch({'action': 'save_person', 'run_id': rid, 'payload': person})


if __name__ == '__main__':
    unittest.main()
