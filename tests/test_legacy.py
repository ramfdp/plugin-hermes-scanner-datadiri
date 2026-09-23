"""Adapted legacy exporter/runtime regressions after moving modules out of root."""
import copy
import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
from docx import Document
from docx.enum.section import WD_ORIENT
from openpyxl import load_workbook
from scanner import exporter, cv_report, tools

ROOT = Path(__file__).resolve().parents[1]
PERSONNEL = {'judul': 'DAFTAR TENAGA AHLI', 'wilayah': 'Data uji', 'pekerjaan': 'Uji ekspor', 'personel': [
    {'nama_personel': 'Contoh 日本 & é', 'nik': '001234', 'pengalaman_kerja_bulan': 37, 'pengalaman_min_kak_tahun': 0},
    {'nama_personel': 'Contoh Nol', 'pengalaman_kerja_bulan': 0, 'pengalaman_kerja_tahun': 0}, {}]}
REVIEW = {'cv_file': 'contoh.pdf', 'biodata': {'nama': 'Contoh 日本 & é', 'masa_kerja': 0},
    'experience_validation': [{'cv_claim': 'Klaim contoh', 'status': 'belum diverifikasi'}],
    'attachment_cross_check': [{'attachment': 'lampiran.pdf', 'field': 'nama', 'cv_value': 'Contoh',
                              'attachment_value': 'Contoh', 'status': 'sesuai'}],
    'findings': ['Perlu bukti tambahan.'], 'internet_sources': [], 'conclusion': 'Perlu review manual.'}


class LegacyTest(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory(); self.addCleanup(temp.cleanup)
        self.root = Path(temp.name)

    def test_excel_identity_zero_calculation_and_layout(self):
        target = self.root / 'nested/out.xlsx'
        self.assertTrue(exporter.export_payload('daftar_tenaga_ahli', PERSONNEL, target)['success'])
        book = load_workbook(target); self.addCleanup(book.close); s = book.active
        for cell, value in {'C6': '001234', 'G6': '0 Tahun', 'H6': 37, 'I6': '3,08', 'H7': 0, 'I7': '0,00', 'I8': None}.items():
            self.assertEqual(s[cell].value, value)
        self.assertEqual(s['C6'].data_type, 's'); self.assertEqual(s['C6'].number_format, '@')
        self.assertEqual(s.freeze_panes, 'A6'); self.assertEqual(s.auto_filter.ref, 'A5:I8')
        self.assertEqual(s.page_setup.orientation, 'landscape'); self.assertEqual(s.page_setup.fitToWidth, 1)
        self.assertEqual(s['A5'].font.name, 'Century Gothic'); self.assertEqual(s['A5'].fill.fgColor.rgb, '008EA9D8')
        self.assertEqual(s.column_dimensions['D'].width, 36)

    def test_explicit_years_empty_personnel_and_long_text(self):
        p = copy.deepcopy(PERSONNEL); p['personel'][0]['pengalaman_kerja_tahun'] = 9.25
        p['personel'][0]['kualifikasi_pendidikan'] = 'Pendidikan contoh cukup panjang ' * 20
        path = self.root / 'a.xlsx'; exporter.export_payload('daftar_tenaga_ahli', p, path)
        book = load_workbook(path)
        self.assertEqual(book.active['I6'].value, '9,25'); self.assertGreater(book.active.row_dimensions[6].height, 54)
        book.close(); p['personel'] = []
        self.assertEqual(exporter.export_payload('daftar_tenaga_ahli', p, path)['row_count'], 0)

    def test_invalid_template_or_payload_does_not_write(self):
        path = self.root / 'bad.xlsx'
        for template, payload in [('unknown', PERSONNEL), ('../escape', PERSONNEL), ('daftar_tenaga_ahli', []),
                                 ('daftar_tenaga_ahli', {}), ('daftar_tenaga_ahli', {**PERSONNEL, 'personel': [None]})]:
            with self.subTest(template=template), self.assertRaises(ValueError):
                exporter.export_payload(template, payload, path)
            self.assertFalse(path.exists())

    def test_bom_and_invalid_json(self):
        path = self.root / 'source.json'; path.write_text(json.dumps(PERSONNEL), encoding='utf-8-sig')
        result = exporter.export_template('daftar_tenaga_ahli', path, self.root / 'a.xlsx')
        self.assertEqual(result['payload_file'], str(path.resolve()))
        for text in ('', '{bad'):
            path.write_text(text)
            with self.assertRaises(ValueError): exporter.load_json(path)
        path.unlink()
        with self.assertRaises(FileNotFoundError): exporter.load_json(path)

    def test_history_and_formula_strings_are_preserved_as_data(self):
        payload = copy.deepcopy(PERSONNEL)
        payload['employment_history'] = [{'nama_personel': 'Contoh', 'employer': '=Example Ltd', 'source_page': 1}]
        result = exporter.export_payload('daftar_tenaga_ahli', payload, self.root / 'a.xlsx')
        self.assertEqual(json.loads(Path(result['payload_file']).read_text(encoding='utf-8')), payload)
        self.assertEqual(result['workbook_data']['Riwayat Pekerjaan'][1][1], '=Example Ltd')
        book = load_workbook(result['output_file']); self.assertEqual(book['Riwayat Pekerjaan']['B2'].data_type, 's'); book.close()

    def test_docx_content_zero_and_extra_workflow_evidence(self):
        p = copy.deepcopy(REVIEW); p['kak_file'] = 'kak.pdf'
        p.update(employer_validation=[{'employer': 'Example', 'status': 'belum dapat diverifikasi', 'evidence': 'Belum tersedia', 'sources': []}],
                 chronology_validation=['Durasi perlu klarifikasi'], source_artifacts={'excel_file': 'cv.xlsx'})
        result = cv_report.render_report(p, self.root / 'a.docx')
        doc = Document(result['output_file'])
        self.assertEqual(doc.paragraphs[0].text, 'LAPORAN REVIEW CV')
        self.assertEqual(doc.tables[0].cell(2, 1).text, '0')
        self.assertEqual(json.loads(Path(result['json_file']).read_text(encoding='utf-8')), p)
        text = '\n'.join(x.text for x in doc.paragraphs) + '\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
        for value in ('Example', 'Belum tersedia', 'cv.xlsx', 'Durasi perlu klarifikasi', 'Contoh 日本 & é'):
            self.assertIn(value, text)

    def test_personnel_word_layout_and_zero(self):
        path = self.root / 'nested/a.docx'
        cv_report.render_personnel_docx({'personel': [{}, {'nik': '000000', 'pengalaman_kerja_bulan': 0,
            'pengalaman_kerja_tahun': None, 'pengalaman_min_kak_tahun': 0}]}, path)
        doc = Document(path)
        self.assertEqual(doc.sections[0].orientation, WD_ORIENT.LANDSCAPE)
        self.assertEqual([c.text for c in doc.tables[0].rows[1].cells], ['1'] + [''] * 8)
        self.assertEqual([c.text for c in doc.tables[0].rows[2].cells], ['2', '', '000000', '', '', '', '0 Tahun', '0', ''])

    def test_empty_and_invalid_docx_payloads(self):
        p = {**REVIEW, 'biodata': 'Belum dipetakan', 'experience_validation': None, 'findings': []}
        result = cv_report.render_report(p, self.root / 'a.docx')
        self.assertEqual(Document(result['output_file']).tables[0].cell(1, 1).text, 'Belum dipetakan')
        for bad in ([], {}, {**REVIEW, 'findings': 'bad'}, {**REVIEW, 'experience_validation': ['bad']},
                    {**REVIEW, 'internet_sources': {}}, {**REVIEW, 'employer_validation': [{'status': 'terverifikasi', 'sources': []}]}):
            with self.subTest(bad=bad), self.assertRaises(ValueError): cv_report.render_report(bad, self.root / 'bad.docx')
        self.assertFalse((self.root / 'bad.docx').exists())

    def test_isolated_child_environment(self):
        original = {'PATH': 'original', 'PYTHONPATH': 'old', 'PYTHONHOME': 'old', 'VIRTUAL_ENV': 'old', 'KEEP': 'value'}
        with patch.dict(os.environ, original, clear=True):
            env = tools._child_env(self.root)
            for key in ('PYTHONPATH', 'PYTHONHOME', 'VIRTUAL_ENV'): self.assertNotIn(key, env)
            self.assertEqual(env['KEEP'], 'value'); self.assertEqual(env['PYTHONIOENCODING'], 'utf-8')
            self.assertEqual(dict(os.environ), original)

    def test_native_cli_clean_package_unicode_and_errors(self):
        shutil.copytree(ROOT / 'scanner', self.root / 'scanner', ignore=shutil.ignore_patterns('__pycache__'))
        self.assertFalse((self.root / 'templates').exists())
        for module, payload, suffix, marker, extra in (
            ('scanner.exporter', PERSONNEL, 'xlsx', 'HERMES_EXPORT_RESULT=', ['--template', 'daftar_tenaga_ahli']),
            ('scanner.cv_report', REVIEW, 'docx', 'HERMES_CV_REPORT_RESULT=', [])):
            path = self.root / ('out.' + suffix)
            command = [sys.executable, '-m', module, *extra, '--payload-json', '-', '--output', str(path)]
            result = subprocess.run(command, input=json.dumps(payload, ensure_ascii=False), text=True, encoding='utf-8',
                capture_output=True, cwd=self.root, env={**os.environ, 'PYTHONIOENCODING': 'cp1252'}, timeout=30)
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            self.assertTrue(json.loads(result.stdout.split(marker)[1])['success'])
            if suffix == 'xlsx':
                book = load_workbook(path); self.assertEqual(book.active['B6'].value, 'Contoh 日本 & é'); book.close()
            else: self.assertEqual(Document(path).tables[0].cell(1, 1).text, 'Contoh 日本 & é')
            path.unlink()
            failed = subprocess.run(command, input='{', text=True, encoding='utf-8', capture_output=True, cwd=self.root, timeout=30)
            self.assertNotEqual(failed.returncode, 0); self.assertFalse(json.loads(failed.stdout.split(marker)[1])['success'])
            self.assertFalse(path.exists())


class RuntimeTest(unittest.TestCase):
    def setUp(self):
        t = tempfile.TemporaryDirectory(); self.addCleanup(t.cleanup); self.root = Path(t.name)
        (self.root / 'scanner').mkdir()
        for module in ('ocr', 'exporter', 'cv_report', 'review'): (self.root / f'scanner/{module}.py').touch()
        (self.root / 'input.pdf').touch()
        for suffix in ('xlsx', 'docx'): (self.root / f'out.{suffix}').write_bytes(b'output')
        env = patch.dict(os.environ, HERMES_SCANNER_PROJECT=str(self.root), HERMES_SCANNER_PYTHON=sys.executable,
                         HERMES_OCR_TIMEOUT_SECONDS='45'); env.start(); self.addCleanup(env.stop)
        self.calls = {'OCR': (tools.scan_document_ocr, {'file_path': str(self.root / 'input.pdf')}),
            'EXPORT': (tools.export_document, {'template': 'daftar_tenaga_ahli', 'payload': {}, 'output_path': str(self.root / 'out.xlsx')}),
            'CV_REPORT': (tools.export_cv_report, {'payload': {}, 'output_path': str(self.root / 'out.docx')})}

    def invoke(self, prefix, stdout=None, returncode=0, exception=None):
        fn, args = self.calls[prefix]
        stdout = stdout if stdout is not None else f'HERMES_{prefix}_RESULT={{"success":true,"ocr_text":"text"}}'
        process = SimpleNamespace(stdout=stdout, stderr='details', returncode=returncode)
        with patch.object(tools.subprocess, 'run', return_value=process, side_effect=exception) as run:
            answer = json.loads(fn(args))
        return answer, run

    def test_safe_subprocess_contract(self):
        for prefix in self.calls:
            result, run = self.invoke(prefix)
            self.assertTrue(result['success']); options = run.call_args.kwargs
            self.assertFalse(options['shell']); self.assertEqual(options['encoding'], 'utf-8')
            self.assertEqual(options['timeout'], 45 if prefix == 'OCR' else 120)
            self.assertEqual(options['cwd'], str(self.root))
            self.assertEqual(run.call_args.args[0][1], '-m')

    def test_nonzero_and_last_marker_never_hide_failures(self):
        for prefix in self.calls:
            result, _ = self.invoke(prefix, returncode=1)
            self.assertFalse(result['success']); self.assertEqual(result['error'], f'{prefix}_PROCESS_FAILED')
        result, _ = self.invoke('OCR', 'HERMES_OCR_RESULT={"success":true,"ocr_text":"old"}\nHERMES_OCR_RESULT={"success":false,"error":"MINERU_OCR_ERROR","output_dir":"partial"}', 1)
        self.assertEqual(result['output_dir'], 'partial'); self.assertEqual(result['error'], 'MINERU_OCR_ERROR')

    def test_invalid_json_shape_and_missing_marker(self):
        for prefix in self.calls:
            for value in ('{', '[]', 'null', '{}', '{"success":"true"}', '{"success":1}'):
                result, _ = self.invoke(prefix, f'HERMES_{prefix}_RESULT={value}')
                self.assertEqual(result['error'], f'{prefix}_INVALID_RESULT')
            result, _ = self.invoke(prefix, 'x' * 6000)
            self.assertEqual(result['error'], f'{prefix}_RESULT_NOT_FOUND'); self.assertLessEqual(len(result['stdout']), 3000)

    def test_empty_text_and_export_file(self):
        for text in (None, '', ' ', []):
            result, _ = self.invoke('OCR', 'HERMES_OCR_RESULT=' + json.dumps({'success': True, 'ocr_text': text}))
            self.assertEqual(result['error'], 'OCR_INVALID_RESULT')
        for prefix, suffix in (('EXPORT', 'xlsx'), ('CV_REPORT', 'docx')):
            path = self.root / ('out.' + suffix); path.unlink()
            for create in (False, True):
                if create: path.touch()
                result, _ = self.invoke(prefix); self.assertEqual(result['error'], f'{prefix}_OUTPUT_NOT_FOUND')

    def test_timeout_and_missing_runtime(self):
        for prefix in self.calls:
            result, _ = self.invoke(prefix, exception=subprocess.TimeoutExpired('test', 1))
            self.assertEqual(result['error'], f'{prefix}_TIMEOUT')
            with patch.dict(os.environ, HERMES_SCANNER_PROJECT=''):
                result, run = self.invoke(prefix); self.assertEqual(result['error'], 'HERMES_SCANNER_PROJECT_NOT_SET'); run.assert_not_called()
            with patch.dict(os.environ, HERMES_SCANNER_PYTHON=str(self.root / 'missing')):
                result, run = self.invoke(prefix); self.assertFalse(result['success']); run.assert_not_called()
        for prefix, module in (('OCR', 'ocr'), ('EXPORT', 'exporter'), ('CV_REPORT', 'cv_report')):
            (self.root / f'scanner/{module}.py').unlink()
            result, run = self.invoke(prefix); self.assertFalse(result['success']); run.assert_not_called()

    def test_bad_arguments_no_process(self):
        with patch.object(tools.subprocess, 'run') as run:
            for fn in (tools.scan_document_ocr, tools.export_document, tools.export_cv_report):
                self.assertFalse(json.loads(fn({}))['success'])
            for prefix in ('EXPORT', 'CV_REPORT'):
                fn, args = self.calls[prefix]; self.assertFalse(json.loads(fn({**args, 'output_path': 'bad.txt'}))['success'])
            for value in ('0', '-1', 'abc'):
                with patch.dict(os.environ, HERMES_OCR_TIMEOUT_SECONDS=value):
                    self.assertEqual(self.invoke('OCR')[0]['error'], 'OCR_PLUGIN_ERROR')
            run.assert_not_called()

    def test_actual_export_tools_current_interpreter(self):
        shutil.rmtree(self.root / 'scanner')
        shutil.copytree(ROOT / 'scanner', self.root / 'scanner', ignore=shutil.ignore_patterns('__pycache__'))
        for prefix, payload in (('EXPORT', PERSONNEL), ('CV_REPORT', REVIEW)):
            fn, args = self.calls[prefix]; Path(args['output_path']).unlink()
            result = json.loads(fn({**args, 'payload': payload}))
            self.assertTrue(result['success'], result); self.assertGreater(Path(result['output_file']).stat().st_size, 100)
