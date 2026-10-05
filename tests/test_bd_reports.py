"""Business-facing export contract, using synthetic data only."""
import copy
import hashlib
import json
import tempfile
import unittest
from pathlib import Path
from openpyxl import load_workbook
import pypdfium2 as pdfium
from scanner.report_labels import AUDIT_LABELS
from scanner.report_presenter import ReportPresenter
from scanner.renderers import pdf, xlsx


def snapshot_fixture():
    ref = {'document_id': 'D001', 'page': 1, 'quote': 'Proyek Jembatan Contoh | PT Contoh | 2020-01 sampai 2021-12'}
    kak_ref = {'document_id': 'D002', 'page': 2, 'quote': 'Pengalaman minimal 12 bulan'}
    history = [{
        'id': 'E001', 'project': 'Proyek Jembatan Contoh', 'employer': 'PT Contoh',
        'role': 'Site Engineer', 'client': 'Dinas Contoh', 'consultant': 'PT Contoh',
        'contractor': 'PT Pelaksana Contoh', 'represented_organization': '',
        'start_date': '2020-01', 'end_date': '2021-12', 'relevant': True,
        'responsibilities': 'Pengawasan pekerjaan struktur',
        'source_refs': [ref], 'supporting_refs': [],
        'supporting_facts': [{'project': 'Proyek Jembatan Contoh', 'role': 'Site Engineer',
                              'start_date': '2020-02', 'source_refs': [ref]}],
    }, {
        'id': 'E002', 'project': 'Proyek Jalan Contoh', 'employer': 'PT Contoh',
        'role': 'Site Engineer', 'start_date': '2021-01', 'end_date': '2021-12',
        'relevant': None, 'responsibilities': 'Pengawasan jalan', 'source_refs': [ref],
        'supporting_refs': [], 'supporting_facts': [],
    }]
    checks = [{
        'code': code, 'applicable': True, 'status': 'memenuhi',
        'finding': 'Data yang diperiksa sesuai.', 'analysis': 'Bukti diperiksa terhadap requirement KAK.',
        'source_refs': [ref], 'kak_refs': [kak_ref], 'computed': True,
    } for code in AUDIT_LABELS]
    indexed = {c['code']: c for c in checks}
    indexed['project_duplicate']['finding'] = 'Tidak ditemukan duplikasi exact maupun kandidat duplikasi pengalaman.'
    indexed['project_overlap'].update(
        status='perlu_klarifikasi', finding='Ditemukan 1 pasangan pengalaman yang bertumpang tindih.',
        analysis='Overlap adalah fakta kronologi, bukan otomatis pelanggaran. Reviewer perlu memastikan apakah pekerjaan memang berjalan bersamaan dan bagaimana KAK memperlakukan periode tersebut.',
        clarification='Konfirmasi proyek yang berjalan bersamaan dan periode yang belum valid; jangan menghapus pengalaman hanya karena overlap.')
    indexed['cv_supporting_document_match'].update(
        status='perlu_klarifikasi',
        finding='Cross-check CV vs dokumen pendukung menemukan 1 konflik dan 40 field/bukti yang belum lengkap.',
        analysis='Cross-check membandingkan field terstruktur tanpa menganggap dokumen yang berbeda sebagai bukti pemalsuan.',
        clarification='Konfirmasi field yang berbeda dan lengkapi bukti.')
    indexed['technical_competency_match'].update(
        applicable=False, status=None, finding='Tidak ada requirement KAK untuk Kesesuaian kompetensi teknis dengan kebutuhan KAK.',
        analysis='Pemeriksaan tidak diterapkan karena KAK tidak memuat requirement terstruktur yang sesuai.')
    indexed['data_anomaly']['status'] = 'perlu_klarifikasi'
    indexed['kak_conclusion'].update(status='perlu_klarifikasi', analysis='Kesimpulan final dihitung deterministik: Tidak Memenuhi mengalahkan Klarifikasi; Klarifikasi mengalahkan Memenuhi. Tidak ada KAK tidak pernah dianggap Memenuhi.')
    person = {
        'id': 'P001', 'name': 'Personel Contoh', 'role': 'Engineer',
        'identity': {'education': 'S1 Teknik Sipil', 'certificate_summary': 'Ahli Teknik Jembatan', 'claimed_months': 24, 'nik': '0000000000000001'},
        'summary': 'Perlu klarifikasi field start_date pada pengalaman E001.',
        'checks': [{'requirement_id': 'R0', 'requirement': 'Lampirkan ijazah', 'status': 'belum_dapat_dinilai',
                    'finding': 'Bukti ijazah belum tersedia.', 'analysis': 'Pendidikan memerlukan bukti.',
                    'clarification': 'Lampirkan ijazah', 'source_refs': [ref]}],
        'audit_checks': checks,
        'document_cross_check': {'details': [{'type': 'experience', 'experience_id': 'E001', 'fields': [
            {'field': 'start_date', 'cv_value': '2020-01', 'support_values': ['2020-02'], 'status': 'conflict'},
            {'field': 'employer', 'cv_value': 'PT Contoh', 'status': 'missing_support'},
        ]}]},
        'anomalies': [{'id': 'A001', 'code': 'cv_support_conflict', 'field': 'start_date', 'material': True,
                       'experience_ids': ['E001'], 'detail': 'E001 field start_date berbeda antara CV dan fakta pendukung.', 'source_refs': [ref]}],
        'final_conclusion': {'status': 'perlu_klarifikasi', 'label': 'Perlu Klarifikasi', 'finding': 'Masih ada informasi yang perlu dikonfirmasi.'},
        'overall': 'belum_dapat_dinilai',
        'kak_match_details': [{'requirement_id': 'R5', 'code': 'experience_duration_match', 'status': 'memenuhi',
                              'requirement_text': 'Pengalaman minimal 12 bulan', 'experience_ids': ['E001'],
                              'finding': 'Pengalaman relevan 24 bulan dan didukung 24 bulan; minimum KAK 12 bulan.',
                              'analysis': 'Durasi memakai union bulan kalender pengalaman yang ditandai relevan; overlap tidak dihitung dua kali.',
                              'source_refs': [ref], 'kak_refs': [kak_ref], 'computed': True}],
        'education_records': [{'level': 'S1', 'degree': 'Sarjana Teknik', 'major': 'Teknik Sipil',
                               'institution': 'Universitas Contoh', 'source_refs': [ref], 'supporting_refs': []}],
        'chronology': {'calendar_months_unique': 24, 'relevant_months': 24, 'supported_relevant_months': 24,
                       'uncertain_entries': [2], 'method': 'Durasi memakai union bulan kalender pengalaman yang ditandai relevan; overlap tidak dihitung dua kali.'},
        'employment_history': history,
        'fact_analysis': {'overlaps': {'pairs': [{'left_id': 'E001', 'right_id': 'E002', 'overlap_start': '2021-01-01',
                            'overlap_end': '2021-12-31', 'calendar_days': 365, 'precision': 'calendar_range_from_partial_dates'}]},
                          'duplicates': {'pairs': []}},
        'employer_checks': [{'employer': 'PT Contoh', 'status': 'belum_dapat_diverifikasi', 'analysis': 'Pemeriksaan daring belum dilakukan.'}],
        'certificates': [{'number': 'CERT-CONTOH-01', 'holder': 'Personel Contoh', 'issuer': 'Penerbit Contoh',
                          'scheme': 'Ahli Teknik Jembatan', 'level': '7', 'issued_on': '2025-01-01', 'expires_on': '2030-01-01',
                          'status': 'belum_dapat_diverifikasi', 'validity': 'belum_melewati_tanggal_akhir',
                          'validity_basis': 'Tanggal akhir dibandingkan dengan tanggal acuan.',
                          'analysis': 'Sertifikat perlu diperiksa pada situs penerbit.', 'reason': 'Pemeriksaan daring belum dilakukan.',
                          'source_refs': [ref]}],
        'certificate_limitation': '', 'findings': ['Konfirmasi start_date pengalaman E001.'],
        'review_note': 'Hasil pemeriksaan tidak menggantikan keputusan pemeriksa.',
    }
    return {
        'schema_version': 2, 'run_id': 'internal-run-id', 'snapshot_sha256': 'a' * 64,
        'project': 'Paket Contoh untuk Pengujian', 'assessment_date': '2026-09-30', 'created_at': '2026-09-30T10:00:00+07:00',
        'documents': [{'id': 'D001', 'kind': 'cv', 'name': 'CV Contoh.pdf', 'sha256': 'b' * 64},
                      {'id': 'D002', 'kind': 'kak', 'name': 'KAK Contoh.pdf', 'sha256': 'c' * 64}],
        'kak': {'title': 'KAK Contoh', 'version': '1', 'package_id': 'Paket Contoh',
                'status': 'dokumen_diberikan_belum_diautentikasi', 'analysis': 'Dokumen KAK tersedia untuk pemeriksaan.'},
        'requirements': [{'id': 'R5', 'role': 'Engineer', 'kind': 'experience', 'minimum_months': 12,
                          'audit_code': 'experience_duration_match', 'text': 'Pengalaman minimal 12 bulan', 'source_refs': [kak_ref]}],
        'people': [person], 'evaluation': {'schema_version': 5, 'criteria_count': 15,
                   'criteria': [{'code': c, 'label': label} for c, label in AUDIT_LABELS.items()], 'implemented_codes': list(AUDIT_LABELS)},
        'limitations': ['P001: audit cv_supporting_document_match memerlukan klarifikasi.', 'D001 halaman 1-2: unreadable'],
        'page_classification': [{'document_id': 'D001', 'first_page': 1, 'last_page': 3, 'kind': 'cv'}],
        'sources': [{'id': 'receipt-internal-01', 'tool': 'browser_snapshot', 'purpose': 'certificate', 'checked_at': '2026-09-30',
                     'evidence_kind': 'page_snapshot', 'success': False, 'urls': ['https://example.org/sertifikat']}],
        'coverage': [], 'report_complete': False,
    }


def pdf_text(path):
    chunks = []
    with pdfium.PdfDocument(str(path)) as document:
        for page in document:
            try:
                text = page.get_textpage()
                try:
                    chunks.append(text.get_text_range())
                finally:
                    text.close()
            finally:
                page.close()
    return '\n'.join(chunks)


class BDPresentationTest(unittest.TestCase):
    def setUp(self):
        self.data = snapshot_fixture()
        self.present = ReportPresenter(self.data)
        self.person = self.data['people'][0]

    def test_all_fifteen_audits_have_business_labels(self):
        self.assertEqual(len(AUDIT_LABELS), 15)
        for check in self.person['audit_checks']:
            self.assertNotIn('_', self.present.check(check)['label'])

    def test_status_mapping_preserves_the_three_final_outcomes(self):
        for internal, label in [('memenuhi', 'Memenuhi'), ('tidak_memenuhi', 'Tidak Memenuhi'), ('perlu_klarifikasi', 'Perlu Klarifikasi')]:
            self.assertEqual(self.present.status(internal), label)

    def test_unknown_status_and_field_do_not_leak_codes(self):
        with self.assertLogs('scanner.report_presenter', 'WARNING'):
            self.assertNotIn('unknown_state', self.present.status('unknown_state'))
            self.assertNotIn('new_field', self.present.field('new_field'))
            self.assertNotIn('new_audit', self.present.audit_label('new_audit'))

    def test_flags_do_not_treat_missing_as_false(self):
        self.assertEqual([self.present.flag(x) for x in [True, False, None]], ['Ya', 'Tidak', 'Belum ditentukan'])

    def test_no_kak_detail_is_not_claimed_to_be_no_requirement(self):
        check = next(c for c in self.person['audit_checks'] if c['code'] == 'technical_competency_match')
        view = self.present.check(check)
        self.assertEqual(view['status'], 'Tidak diterapkan')
        self.assertIn('tidak menyatakan bahwa KAK pasti tidak mensyaratkannya', view['analysis'])
        self.assertFalse(check['applicable'])
        self.assertIsNone(check['status'])

    def test_missing_information_count_is_not_a_document_count(self):
        check = next(c for c in self.person['audit_checks'] if c['code'] == 'cv_supporting_document_match')
        text = self.present.check(check)['finding']
        self.assertIn('1 perbedaan', text)
        self.assertIn('40 butir informasi', text)
        self.assertIn('bukan jumlah dokumen', text)
        self.assertNotIn('40 dokumen', text)

    def test_experience_names_are_person_scoped(self):
        another = copy.deepcopy(self.person)
        another['employment_history'][0]['project'] = 'Proyek Personel Kedua'
        self.assertIn('Proyek Jembatan Contoh', self.present.experience(self.person, 'E001'))
        self.assertIn('Proyek Personel Kedua', self.present.experience(another, 'E001'))

    def test_missing_experience_does_not_fabricate_a_project(self):
        with self.assertLogs('scanner.report_presenter', 'WARNING'):
            self.assertIn('belum teridentifikasi', self.present.experience(self.person, 'E999'))

    def test_source_names_and_quotes_are_verbatim(self):
        literal = 'Cross-check project_duplicate: field start_date & <skema>'
        self.data['documents'][0]['name'] = 'Project_Overlap & Field.pdf'
        present = ReportPresenter(self.data)
        text = present.references([{'document_id': 'D001', 'page': 7, 'quote': literal}])
        self.assertEqual(text, 'Project_Overlap & Field.pdf, halaman 7: ' + literal)

    def test_narrative_protects_project_titles(self):
        title = 'Exact Match Project_Overlap Study'
        self.person['employment_history'][0]['project'] = title
        present = ReportPresenter(self.data)
        text = present.narrative(f'Cross-check field start_date pada {title} untuk E001.', self.person)
        self.assertIn(title, text)
        self.assertNotIn('field start_date', text)
        self.assertNotIn('E001', text)

    def test_quoted_engine_sentence_is_not_rewritten_inside_narrative(self):
        quote = 'Tidak ditemukan duplikasi exact maupun kandidat duplikasi pengalaman.'
        self.person['employment_history'][0]['source_refs'][0]['quote'] = quote
        present = ReportPresenter(self.data)
        self.assertIn(quote, present.narrative('Kutipan sumber: "' + quote + '"', self.person))

    def test_inserted_requirement_text_is_not_translated_again(self):
        self.data['requirements'][0]['text'] = 'Exact role E001 project_duplicate'
        present = ReportPresenter(self.data)
        result = present.narrative('Bukti untuk requirement R5.', self.person)
        self.assertIn('Exact role E001 project_duplicate', result)
        self.assertNotIn('R5', result)

    def test_uncertain_periods_resolve_to_the_correct_project(self):
        self.assertIn('Proyek Jalan Contoh', self.present.uncertain_experiences(self.person))
        self.assertNotIn('Proyek Jembatan Contoh', self.present.uncertain_experiences(self.person))

    def test_short_source_values_do_not_split_internal_identifiers(self):
        self.person['name'] = 'a'
        self.person['certificates'][0]['number'] = '001'
        present = ReportPresenter(self.data)
        result = present.narrative('audit project_duplicate pada E001', self.person)
        self.assertNotIn('project_duplicate', result)
        self.assertNotIn('E001', result)
        self.assertIn('Proyek Jembatan Contoh', result)

    def test_source_reference_contains_document_name_and_page(self):
        self.assertIn('CV Contoh.pdf, halaman 1:', self.present.references(self.person['employment_history'][0]['source_refs']))

    def test_labels_keep_portal_identity_separate_from_date_validity(self):
        cert = self.person['certificates'][0]
        self.assertEqual(self.present.status(cert['status']), 'Belum dapat diverifikasi')
        self.assertIn('tanggal akhir', self.present.status(cert['validity']))


class BDExportTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.data = snapshot_fixture()

    def export(self):
        pdf_path, xlsx_path = self.root / 'report.pdf', self.root / 'report.xlsx'
        pdf.render(self.data, pdf_path)
        xlsx.render(self.data, xlsx_path)
        return pdf_path, xlsx_path

    def test_export_does_not_mutate_snapshot_or_evaluation(self):
        before = copy.deepcopy(self.data)
        digest = hashlib.sha256(json.dumps(self.data, sort_keys=True).encode()).hexdigest()
        self.export()
        self.assertEqual(self.data, before)
        self.assertEqual(hashlib.sha256(json.dumps(self.data, sort_keys=True).encode()).hexdigest(), digest)

    def test_both_formats_share_labels_statuses_and_evidence(self):
        pdf_path, xlsx_path = self.export()
        text = ' '.join(pdf_text(pdf_path).split())
        book = load_workbook(xlsx_path, data_only=True)
        self.addCleanup(book.close)
        self.assertEqual(book['Ringkasan Personel'].max_row, len(self.data['people']) + 5)
        self.assertEqual(book['Ringkasan Pemeriksaan']['D2'].value, 'Perlu Klarifikasi')
        audit = book['Hasil Pemeriksaan']
        rows = list(audit.iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(rows), 15)
        for row in rows:
            self.assertIn(row[3], text)
            self.assertIn(row[4], text)
        self.assertIn('CV Contoh.pdf, halaman 1:', text)
        self.assertTrue(any('KAK Contoh.pdf, halaman 2:' in str(row[-1]) for row in rows))
        self.assertIn('40 butir informasi', text)

    def test_internal_identifiers_are_not_in_business_exports(self):
        pdf_path, xlsx_path = self.export()
        text = pdf_text(pdf_path)
        book = load_workbook(xlsx_path, data_only=True)
        self.addCleanup(book.close)
        cells = '\n'.join(str(v) for s in book for row in s.iter_rows(values_only=True) for v in row if v is not None)
        for token in [*AUDIT_LABELS, 'missing_support', 'cv_support_conflict', 'calendar_range_from_partial_dates',
                      'start_date', 'E001', 'E002', 'P001', 'D001', 'receipt-internal-01', 'browser_snapshot',
                      'Status legacy', 'Status final internal', 'Schema review', 'Snapshot SHA256', 'kak_match_details']:
            with self.subTest(token=token):
                self.assertNotIn(token, text)
                self.assertNotIn(token, cells)
        self.assertTrue(all(s.sheet_state == 'visible' for s in book))
        self.assertNotIn('Kode Audit', cells)

    def test_literal_source_with_technical_words_is_not_rewritten(self):
        literal = 'project_duplicate & start_date are printed on the original.'
        self.data['people'][0]['employment_history'][0]['source_refs'][0]['quote'] = literal
        pdf_path, xlsx_path = self.export()
        self.assertIn(literal, pdf_text(pdf_path))
        book = load_workbook(xlsx_path)
        self.addCleanup(book.close)
        self.assertTrue(any(literal in str(v) for sheet in book for row in sheet.iter_rows(values_only=True) for v in row))

    def test_workbook_preserves_nik_and_prevents_formula_injection(self):
        self.data['people'][0]['name'] = '=HYPERLINK("https://example.org","test")'
        _, xlsx_path = self.export()
        book = load_workbook(xlsx_path, data_only=False)
        self.addCleanup(book.close)
        self.assertEqual(book['Ringkasan Personel']['C6'].value, '0000000000000001')
        self.assertEqual(book['Ringkasan Personel']['B6'].data_type, 's')
        self.assertFalse(any(c.data_type == 'f' for s in book for row in s for c in row))

    def test_partial_dates_remain_explicit(self):
        pdf_path, xlsx_path = self.export()
        self.assertIn('Perkiraan rentang; tanggal belum lengkap', pdf_text(pdf_path))
        book = load_workbook(xlsx_path)
        self.addCleanup(book.close)
        rows = list(book['Periode dan Pencatatan Ganda'].iter_rows(min_row=2, values_only=True))
        self.assertIn('Perkiraan rentang; tanggal belum lengkap', rows[0])
        self.assertIn(365, rows[0])

    def test_nonapplicable_and_empty_sections_render_without_inventing_success(self):
        person = self.data['people'][0]
        person['audit_checks'] = []
        person['checks'] = []
        person['kak_match_details'] = []
        person['employment_history'] = []
        person['document_cross_check'] = {}
        person['anomalies'] = []
        person['fact_analysis'] = {}
        person['certificates'] = []
        person['certificate_limitation'] = 'Sertifikat belum tersedia.'
        self.data['requirements'] = []
        self.data['kak']['status'] = 'tidak_tersedia'
        pdf_path, xlsx_path = self.export()
        self.assertIn('Kesesuaian belum dapat dinilai', pdf_text(pdf_path))
        book = load_workbook(xlsx_path)
        self.addCleanup(book.close)
        self.assertEqual(book['Ringkasan Pemeriksaan']['D2'].value, 'Perlu Klarifikasi')


if __name__ == '__main__':
    unittest.main()
