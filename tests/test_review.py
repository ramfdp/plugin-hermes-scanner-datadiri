import copy
import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from openpyxl import load_workbook
from scanner import review, storage
from scanner.validation import chronology, certificate_result
from helpers import create_run, plan_payload, person_payload, complete_run, make_receipt, CERT_PAGES


class ReviewTest(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.root = Path(temp.name)
        env = patch.dict(os.environ, HERMES_SCANNER_OUTPUT_DIR=str(self.root / 'out'))
        env.start()
        self.addCleanup(env.stop)

    def test_two_real_artifacts_share_snapshot(self):
        rid = complete_run(self.root)
        result = review.dispatch({'action': 'export', 'run_id': rid})
        self.assertTrue(result['success'], result)
        self.assertEqual(len(result['artifacts']), 2)
        self.assertFalse(result['report_complete'])
        xlsx, pdf = [Path(a['path']) for a in result['artifacts']]
        book = load_workbook(xlsx, read_only=True, data_only=True)
        self.assertEqual(book['Ringkasan Personel'].max_row, 7)
        self.assertEqual(book['Verifikasi Sertifikat']['J2'].value, 'terverifikasi')
        self.assertEqual(book['Verifikasi Sertifikat']['J3'].value, 'belum_dapat_diverifikasi')
        self.assertEqual(book['Info Pemeriksaan']['B3'].value, result['snapshot_sha256'])
        book.close()
        import pypdfium2 as pdfium
        with pdfium.PdfDocument(str(pdf)) as document:
            self.assertGreaterEqual(len(document), 4)
            chunks = []
            for index in range(len(document)):
                page = document[index]
                textpage = page.get_textpage()
                chunks.append(textpage.get_text_range())
                textpage.close()
                page.close()
        text = '\n'.join(chunks)
        for term in ('Personel Contoh A', 'Personel Contoh B', 'Lampirkan ijazah', result['snapshot_sha256'][:16]):
            self.assertIn(term, text)
        self.assertNotIn('.json', result['chat_markdown'])

    def test_pdf_failure_keeps_only_valid_excel(self):
        rid = complete_run(self.root)
        with patch('scanner.renderers.pdf.render', side_effect=RuntimeError('Synthetic failure')):
            result = review.dispatch({'action': 'export', 'run_id': rid})
        self.assertFalse(result['success'])
        self.assertEqual([a['name'] for a in result['artifacts']], ['Ringkasan_Tenaga_Ahli.xlsx'])
        self.assertEqual(result['errors'][0]['file'], 'Laporan_Verifikasi_CV_KAK.pdf')

    def test_repeat_export_never_overwrites_prior_files(self):
        rid = complete_run(self.root)
        a = review.dispatch({'action': 'export', 'run_id': rid})
        b = review.dispatch({'action': 'export', 'run_id': rid})
        self.assertNotEqual(a['artifacts'][0]['path'], b['artifacts'][0]['path'])
        self.assertTrue(Path(a['artifacts'][0]['path']).is_file())

    def test_missing_person_blocks_export(self):
        rid = create_run(self.root)
        review.dispatch({'action': 'plan', 'run_id': rid, 'payload': plan_payload()})
        review.dispatch({'action': 'save_person', 'run_id': rid, 'payload': person_payload(0)})
        with self.assertRaisesRegex(ValueError, 'P2'):
            review.dispatch({'action': 'export', 'run_id': rid})

    def test_expected_count_prevents_lost_cv(self):
        rid = create_run(self.root, expected=3)
        with self.assertRaisesRegex(ValueError, 'expected_person_count'):
            review.dispatch({'action': 'plan', 'run_id': rid, 'payload': plan_payload()})

    def test_all_cv_pages_require_coverage(self):
        rid = create_run(self.root)
        plan = plan_payload()
        plan['coverage'] = plan['coverage'][1:]
        with self.assertRaisesRegex(ValueError, 'halaman'):
            review.dispatch({'action': 'plan', 'run_id': rid, 'payload': plan})

    def test_unreadable_pages_are_explicit_limitations(self):
        rid = create_run(self.root)
        plan = plan_payload()
        plan['roster'][1]['certificate_inventory'] = []
        plan['coverage'][-1] = {'document_id': 'D003', 'first_page': 2, 'last_page': 2, 'person_ids': [], 'reason': 'unreadable'}
        review.dispatch({'action': 'plan', 'run_id': rid, 'payload': plan})
        for i in range(2):
            person = person_payload(i)
            if i == 1:
                person['certificates'] = []
                person['certificate_limitation'] = 'Lampiran sertifikat tidak terbaca.'
            review.dispatch({'action': 'save_person', 'run_id': rid, 'payload': person})
        data = review.snapshot(storage.run_dir(rid))
        self.assertTrue(any('unreadable' in x for x in data['limitations']))

    def test_missing_criteria_or_invented_quote_is_rejected(self):
        rid = create_run(self.root)
        review.dispatch({'action': 'plan', 'run_id': rid, 'payload': plan_payload()})
        for mutation in ('omit', 'invent'):
            person = person_payload(0)
            if mutation == 'omit':
                person['checks'] = []
            else:
                person['checks'][0]['source_refs'][0]['quote'] = 'Invented education evidence'
            with self.subTest(mutation=mutation), self.assertRaises(ValueError):
                review.dispatch({'action': 'save_person', 'run_id': rid, 'payload': person})

    def test_missing_inventory_certificate_is_rejected(self):
        rid = create_run(self.root)
        review.dispatch({'action': 'plan', 'run_id': rid, 'payload': plan_payload()})
        person = person_payload(0)
        person['certificates'] = []
        with self.assertRaisesRegex(ValueError, 'certificate_inventory'):
            review.dispatch({'action': 'save_person', 'run_id': rid, 'payload': person})

    def test_no_kak_does_not_mean_zero_requirements_met(self):
        rid = create_run(self.root)
        plan = plan_payload()
        plan['requirements'] = []
        plan['kak']['analysis'] = 'KAK tidak dapat digunakan; persyaratan belum dapat dinilai pada pengujian ini.'
        review.dispatch({'action': 'plan', 'run_id': rid, 'payload': plan})
        person = person_payload(0)
        person['checks'] = []
        result = review.dispatch({'action': 'save_person', 'run_id': rid, 'payload': person})
        self.assertEqual(result['overall'], 'tidak_ada_kriteria_KAK')

    def test_search_result_is_not_verification_evidence(self):
        rid = create_run(self.root)
        review.dispatch({'action': 'plan', 'run_id': rid, 'payload': plan_payload()})
        receipt = make_receipt(rid, tool='web_search')
        person = person_payload(0)
        person['certificates'][0].update(receipt_id=receipt['receipt_id'], quote=CERT_PAGES[0])
        result = review.dispatch({'action': 'save_person', 'run_id': rid, 'payload': person})
        self.assertEqual(result['certificates'][0]['status'], 'belum_dapat_diverifikasi')

    def test_quote_cannot_mix_unrelated_records_or_invent_receipt(self):
        rid = create_run(self.root)
        review.dispatch({'action': 'plan', 'run_id': rid, 'payload': plan_payload()})
        receipt = make_receipt(rid, content='TEST-001 Personel Contoh B Skema Contoh Penerbit Contoh')
        person = person_payload(0)
        person['certificates'][0].update(receipt_id=receipt['receipt_id'], quote='TEST-001 Personel Contoh B Skema Contoh Penerbit Contoh')
        result = review.dispatch({'action': 'save_person', 'run_id': rid, 'payload': person})
        self.assertNotEqual(result['certificates'][0]['status'], 'terverifikasi')
        person['certificates'][0]['receipt_id'] = '../secret'
        with self.assertRaises(ValueError):
            review.dispatch({'action': 'save_person', 'run_id': rid, 'payload': person})

    def test_certificate_identity_and_expiry_are_separate(self):
        rid = create_run(self.root)
        review.dispatch({'action': 'plan', 'run_id': rid, 'payload': plan_payload()})
        receipt = make_receipt(rid)
        person = person_payload(0)
        person['certificates'][0].update(receipt_id=receipt['receipt_id'], quote=CERT_PAGES[0], expires_on='2025-12-31')
        review.dispatch({'action': 'save_person', 'run_id': rid, 'payload': person})
        stored = storage.load(storage.run_dir(rid)/'people/P1.json')['certificates'][0]
        self.assertEqual(stored['status'], 'terverifikasi')
        self.assertEqual(stored['validity'], 'kedaluwarsa')

    def test_history_union_and_missing_precision(self):
        result = chronology([{'start_date': '2020-01', 'end_date': '2020-12', 'relevant': True},
                             {'start_date': '2020-07', 'end_date': '2021-06', 'relevant': True},
                             {'start_date': '2019', 'end_date': '2020', 'relevant': True}])
        self.assertEqual(result['calendar_months_unique'], 18)
        self.assertEqual(result['uncertain_entries'], [3])
        with self.assertRaises(ValueError):
            chronology([{'start_date': '2022-01', 'end_date': '2020-01'}])

    def test_read_pagination_and_cached_ocr(self):
        rid = create_run(self.root)
        with patch('scanner.ocr.run_mineru') as runner:
            self.assertTrue(review.dispatch({'action': 'document', 'run_id': rid, 'payload': {'document_id': 'D001'}})['cached'])
            runner.assert_not_called()
        result = review.dispatch({'action': 'read_document', 'run_id': rid, 'payload': {'document_id': 'D001'}})
        self.assertEqual(result['next']['page'], 2)
        final = review.dispatch({'action': 'read_document', 'run_id': rid, 'payload': result['next']})
        self.assertIsNone(final['next'])
        self.assertIn('Personel Contoh B', final['text'])

    def test_plan_cannot_change_after_people_are_saved(self):
        rid = complete_run(self.root)
        with self.assertRaisesRegex(ValueError, 'run baru'):
            review.dispatch({'action': 'plan', 'run_id': rid, 'payload': plan_payload()})

    def test_run_path_cannot_escape_root(self):
        for value in ('../outside', '/tmp', None, 'x'*32):
            with self.subTest(value=value), self.assertRaises(ValueError):
                storage.run_dir(value)

    def test_contract_is_available_without_model(self):
        result = review.dispatch({'action': 'help'})
        for term in ('certificate_inventory', 'source_refs', 'save_person', 'read_document'):
            self.assertIn(term, result['format'])

class EvidenceOwnershipTest(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory(); self.addCleanup(temp.cleanup)
        self.root = Path(temp.name)
        env = patch.dict(os.environ, HERMES_SCANNER_OUTPUT_DIR=str(self.root / 'output')); env.start(); self.addCleanup(env.stop)
        self.rid = create_run(self.root)
        review.dispatch({'action': 'plan', 'run_id': self.rid, 'payload': plan_payload()})

    def test_other_person_page_is_rejected(self):
        p = person_payload(0)
        p['checks'][0]['source_refs'][0]['page'] = 2
        with self.assertRaisesRegex(ValueError, 'bukan milik'):
            review.dispatch({'action': 'save_person', 'run_id': self.rid, 'payload': p})

    def test_certificate_values_cannot_be_invented_from_web(self):
        p = person_payload(0); p['certificates'][0]['number'] = 'INVENTED-003'
        with self.assertRaisesRegex(ValueError, 'tidak didukung'):
            review.dispatch({'action': 'save_person', 'run_id': self.rid, 'payload': p})

    def test_future_months_are_not_experience_already_earned(self):
        from datetime import date
        from scanner.validation import chronology
        value = chronology([{'start_date': '2026-08', 'end_date': '2028-01', 'relevant': True}], date(2026, 9, 23))
        self.assertEqual(value['calendar_months_unique'], 1)
        self.assertEqual(value['uncertain_entries'], [1])
