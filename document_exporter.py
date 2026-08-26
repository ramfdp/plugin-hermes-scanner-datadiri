from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from docx import Document


FIELD_LABELS = {
    "provinsi": "Provinsi",
    "kabupaten_kota": "Kabupaten/Kota",
    "nik": "NIK",
    "nama": "Nama",
    "tempat_tanggal_lahir": "Tempat/Tanggal Lahir",
    "jenis_kelamin": "Jenis Kelamin",
    "golongan_darah": "Golongan Darah",
    "alamat": "Alamat",
    "rt_rw": "RT/RW",
    "kelurahan_desa": "Kelurahan/Desa",
    "kecamatan": "Kecamatan",
    "agama": "Agama",
    "status_perkawinan": "Status Perkawinan",
    "pekerjaan": "Pekerjaan",
    "kewarganegaraan": "Kewarganegaraan",
}


def export_ktp_excel(data: dict, output_path):
    output_path = Path(output_path)

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True
    )

    workbook = Workbook()

    sheet = workbook.active
    sheet.title = "Data KTP"

    sheet["A1"] = "HASIL EKSTRAKSI DATA KTP"
    sheet["A1"].font = Font(
        bold=True,
        size=14
    )

    sheet.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=2
    )

    row = 3

    for key, label in FIELD_LABELS.items():
        sheet.cell(
            row=row,
            column=1,
            value=label
        )

        sheet.cell(
            row=row,
            column=2,
            value=data.get(key, "")
        )

        sheet.cell(
            row=row,
            column=1
        ).font = Font(bold=True)

        sheet.cell(
            row=row,
            column=1
        ).alignment = Alignment(
            vertical="top"
        )

        sheet.cell(
            row=row,
            column=2
        ).alignment = Alignment(
            wrap_text=True,
            vertical="top"
        )

        row += 1

    sheet.column_dimensions[
        get_column_letter(1)
    ].width = 25

    sheet.column_dimensions[
        get_column_letter(2)
    ].width = 55

    workbook.save(output_path)

    return str(output_path)


def export_ktp_word(data: dict, output_path):
    output_path = Path(output_path)

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True
    )

    document = Document()

    document.add_heading(
        "Hasil Ekstraksi Data KTP",
        level=1
    )

    table = document.add_table(
        rows=1,
        cols=2
    )

    table.style = "Table Grid"

    header = table.rows[0].cells

    header[0].text = "Field"
    header[1].text = "Nilai"

    for key, label in FIELD_LABELS.items():
        cells = table.add_row().cells

        cells[0].text = label
        cells[1].text = str(
            data.get(key, "")
        )

    document.save(output_path)

    return str(output_path)