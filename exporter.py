import argparse
import json
import math
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins

TEMPLATE_ID = "daftar_tenaga_ahli"
REQUIRED_FIELDS = ("judul", "wilayah", "pekerjaan", "personel")


def load_json(path: Path) -> dict:
    if not path.exists():
        raise FileNotFoundError(f"File JSON tidak ditemukan: {path}")
    if path.stat().st_size == 0:
        raise ValueError(f"File JSON kosong: {path}")
    try:
        return json.loads(path.read_text(encoding="utf-8-sig"))
    except json.JSONDecodeError as exc:
        raise ValueError(
            f"JSON tidak valid pada '{path}' baris {exc.lineno}, kolom {exc.colno}: {exc.msg}"
        ) from exc


def truncate_decimal(value: float, decimals: int = 2) -> float:
    multiplier = 10 ** decimals
    return math.floor(value * multiplier) / multiplier


def calculate_years(months) -> float:
    try:
        months = float(months)
    except (TypeError, ValueError):
        return 0.0
    return truncate_decimal(months / 12, 2)


def format_indonesian_decimal(value) -> str:
    try:
        return f"{float(value):.2f}".replace(".", ",")
    except (TypeError, ValueError):
        return ""


def validate_payload(payload: dict, schema: dict | None = None):
    if not isinstance(payload, dict):
        raise ValueError("Payload harus berupa object.")
    required = REQUIRED_FIELDS if schema is None else schema.get("required_fields", [])
    missing = [field for field in required if field not in payload]
    if missing:
        raise ValueError("Field wajib belum tersedia: " + ", ".join(missing))
    personel = payload.get("personel")
    if not isinstance(personel, list):
        raise ValueError("'personel' harus berupa array/list.")
    if any(not isinstance(person, dict) for person in personel):
        raise ValueError("Setiap item 'personel' harus berupa object.")


def render_daftar_tenaga_ahli(payload: dict, output_path: Path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Daftar Tenaga Ahli"
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_margins = PageMargins(left=0.25, right=0.25, top=0.35, bottom=0.35, header=0.15, footer=0.15)

    font_name = "Century Gothic"
    centered = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_aligned = Alignment(horizontal="left", vertical="center", wrap_text=True)
    body_font = Font(name=font_name, size=8)
    header_font = Font(name=font_name, size=8, bold=True, color="000000")
    header_fill = PatternFill(fill_type="solid", fgColor="8EA9D8")
    side = Side(style="thin", color="000000")
    border = Border(left=side, right=side, top=side, bottom=side)

    titles = [payload.get("judul", "DAFTAR TENAGA AHLI"), payload.get("wilayah", ""), payload.get("pekerjaan", "")]
    for row, title in enumerate(titles, 1):
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
        cell = sheet.cell(row, 1, title)
        cell.font = Font(name=font_name, size=11 if row == 1 else 10, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.row_dimensions[row].height = 22 if row == 1 else 21
    sheet.row_dimensions[4].height = 10

    headers = ["No", "Nama Personel", "NIK", "Jabatan Personel", "Kualifikasi\nPendidikan",
               "Sertifikat Keahlian", "Pengalaman min\ndalam KAK\n(Tahun)",
               "Pengalaman\nKerja (Bulan)", "Pengalaman\nKerja (Tahun)"]
    for column, header in enumerate(headers, 1):
        cell = sheet.cell(5, column, header)
        cell.font, cell.fill, cell.border, cell.alignment = header_font, header_fill, border, centered
    sheet.row_dimensions[5].height = 48
    for column, width in zip("ABCDEFGHI", (5, 18, 24, 36, 16, 23, 19, 17, 17)):
        sheet.column_dimensions[column].width = width

    personel = payload.get("personel", [])
    for number, person in enumerate(personel, 1):
        row = number + 5
        months = person.get("pengalaman_kerja_bulan", "")
        years = person.get("pengalaman_kerja_tahun")
        if years in (None, ""):
            years = calculate_years(months) if months not in (None, "") else ""
        min_kak = person.get("pengalaman_min_kak_tahun", "")
        values = [number, person.get("nama_personel", ""), str(person.get("nik", "")),
                  person.get("jabatan_personel", ""), person.get("kualifikasi_pendidikan", ""),
                  person.get("sertifikat_keahlian", ""),
                  f"{min_kak} Tahun" if min_kak not in (None, "") else "", months,
                  format_indonesian_decimal(years) if years != "" else ""]
        for column, value in enumerate(values, 1):
            cell = sheet.cell(row, column, value)
            cell.font, cell.border = body_font, border
            cell.alignment = left_aligned if column in (2, 3) else centered
        sheet.cell(row, 3).number_format = "@"
        if isinstance(months, (int, float)):
            sheet.cell(row, 8).number_format = "0"
        sheet.row_dimensions[row].height = 54

    if personel:
        last_row = 5 + len(personel)
        sheet.print_area = f"A1:I{last_row}"
        sheet.auto_filter.ref = f"A5:I{last_row}"
    sheet.freeze_panes = "A6"
    sheet.print_options.horizontalCentered = True
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def export_template(template_id: str, payload_path: Path, output_path: Path):
    result = export_payload(template_id, load_json(payload_path), output_path)
    result["payload_file"] = str(payload_path.resolve())
    return result


def export_payload(template_id: str, payload: dict, output_path: Path):
    if template_id != TEMPLATE_ID:
        raise ValueError(f"Renderer untuk template '{template_id}' belum tersedia.")
    validate_payload(payload)
    render_daftar_tenaga_ahli(payload, output_path)
    return {"success": True, "template": template_id,
            "output_file": str(output_path.resolve()), "row_count": len(payload["personel"])}


def main():
    parser = argparse.ArgumentParser(description="Hermes document template exporter")
    parser.add_argument("--template", required=True, help="Template ID: daftar_tenaga_ahli")
    group = parser.add_mutually_exclusive_group(required=True)
    group.add_argument("--payload", type=Path, help="Path file JSON hasil mapping Hermes")
    group.add_argument("--payload-json", help="JSON hasil mapping Hermes, atau '-' untuk stdin")
    parser.add_argument("--output", required=True, type=Path, help="Path output Excel .xlsx")
    args = parser.parse_args()
    for stream in (sys.stdin, sys.stdout):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8")
    try:
        if args.payload_json is not None:
            payload_text = sys.stdin.read() if args.payload_json == "-" else args.payload_json
            result = export_payload(args.template, json.loads(payload_text), args.output)
        else:
            result = export_template(args.template, args.payload, args.output)
    except Exception as exc:
        result = {"success": False, "error": type(exc).__name__, "message": str(exc)}
    print("HERMES_EXPORT_RESULT=" + json.dumps(result, ensure_ascii=False))
    if not result["success"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
