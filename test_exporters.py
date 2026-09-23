import copy
import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from openpyxl import load_workbook

import cv_report
import exporter


PERSONNEL = {
    "judul": "DAFTAR TENAGA AHLI", "wilayah": "Data uji", "pekerjaan": "Uji ekspor",
    "personel": [{"nama_personel": "Contoh & Uji <A>", "nik": "0012345678901234",
                  "pengalaman_kerja_bulan": 37, "pengalaman_min_kak_tahun": 0},
                 {"nama_personel": "Contoh Nol", "pengalaman_kerja_bulan": 0,
                  "pengalaman_kerja_tahun": 0}, {}],
}
REVIEW = {
    "cv_file": "contoh.pdf", "biodata": {"nama": "Contoh & Uji <A>", "masa_kerja": 0},
    "experience_validation": [{"cv_claim": "Klaim contoh", "status": "belum diverifikasi"}],
    "attachment_cross_check": [{"attachment": "lampiran.pdf", "field": "nama",
                                "cv_value": "Contoh", "attachment_value": "Contoh", "status": "sesuai"}],
    "findings": ["Perlu bukti tambahan."], "internet_sources": [], "conclusion": "Perlu review manual.",
}


class ExporterTest(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.root = Path(temp.name)

    def test_clean_install_excel_and_layout(self):
        target = self.root / "nested" / "personnel.xlsx"
        result = exporter.export_payload("daftar_tenaga_ahli", copy.deepcopy(PERSONNEL), target)
        self.assertTrue(result["success"])
        self.assertEqual(result["row_count"], 3)
        workbook = load_workbook(target)
        self.addCleanup(workbook.close)
        sheet = workbook.active
        self.assertEqual(sheet["C6"].value, "0012345678901234")
        self.assertEqual(sheet["C6"].data_type, "s")
        self.assertEqual(sheet["C6"].number_format, "@")
        self.assertEqual(sheet["G6"].value, "0 Tahun")
        self.assertEqual(sheet["H6"].value, 37)
        self.assertEqual(sheet["I6"].value, "3,08")
        self.assertEqual(sheet["H7"].value, 0)
        self.assertEqual(sheet["I7"].value, "0,00")
        self.assertEqual(sheet["I8"].value, None)
        self.assertEqual(sheet.freeze_panes, "A6")
        self.assertEqual(sheet.auto_filter.ref, "A5:I8")
        self.assertEqual(sheet.page_setup.orientation, "landscape")
        self.assertEqual(str(sheet.page_setup.paperSize), str(sheet.PAPERSIZE_A4))
        self.assertEqual(sheet.page_setup.fitToWidth, 1)
        self.assertEqual(sheet.page_setup.fitToHeight, 0)
        self.assertEqual(sheet["A5"].font.name, "Century Gothic")
        self.assertEqual(sheet["A5"].fill.fgColor.rgb, "008EA9D8")
        self.assertEqual(sheet["B6"].alignment.horizontal, "left")
        self.assertEqual(sheet["D6"].alignment.horizontal, "center")
        self.assertEqual(sheet.row_dimensions[6].height, 54)
        self.assertEqual(sheet.column_dimensions["D"].width, 36)

    def test_explicit_years_and_empty_personnel(self):
        payload = copy.deepcopy(PERSONNEL)
        payload["personel"][0]["pengalaman_kerja_tahun"] = 9.25
        target = self.root / "explicit.xlsx"
        exporter.export_payload("daftar_tenaga_ahli", payload, target)
        workbook = load_workbook(target)
        self.assertEqual(workbook.active["I6"].value, "9,25")
        workbook.close()
        payload["personel"] = []
        result = exporter.export_payload("daftar_tenaga_ahli", payload, self.root / "empty.xlsx")
        self.assertEqual(result["row_count"], 0)

    def test_invalid_excel_payload_and_template_do_not_write(self):
        target = self.root / "invalid.xlsx"
        cases = [("unknown", PERSONNEL), ("../other", PERSONNEL)]
        for payload in ([], {}, {**PERSONNEL, "personel": {}}, {**PERSONNEL, "personel": [None]}):
            cases.append(("daftar_tenaga_ahli", payload))
        for template, payload in cases:
            with self.subTest(template=template, payload=payload), self.assertRaises(ValueError):
                exporter.export_payload(template, payload, target)
            self.assertFalse(target.exists())

    def test_file_payload_supports_bom_and_errors(self):
        source = self.root / "payload.json"
        source.write_text(json.dumps(PERSONNEL, ensure_ascii=False), encoding="utf-8-sig")
        result = exporter.export_template("daftar_tenaga_ahli", source, self.root / "file.xlsx")
        self.assertEqual(result["payload_file"], str(source.resolve()))
        for content, message in (("", "kosong"), ("{bad", "JSON tidak valid")):
            source.write_text(content, encoding="utf-8")
            with self.assertRaisesRegex(ValueError, message):
                exporter.load_json(source)
        source.unlink()
        with self.assertRaises(FileNotFoundError):
            exporter.load_json(source)

    def test_review_content_and_zero_values(self):
        target = self.root / "nested" / "review.docx"
        payload = copy.deepcopy(REVIEW)
        payload["kak_file"] = "kak.pdf"
        result = cv_report.render_report(payload, target)
        document = Document(target)
        self.assertEqual(result["finding_count"], 1)
        self.assertEqual(document.paragraphs[0].text, "LAPORAN REVIEW CV")
        self.assertEqual(len([p for p in document.paragraphs if p.style.name == "Heading 1"]), 6)
        self.assertIn("File KAK: kak.pdf", [p.text for p in document.paragraphs])
        self.assertEqual(document.tables[0].cell(1, 1).text, "Contoh & Uji <A>")
        self.assertEqual(document.tables[0].cell(2, 1).text, "0")
        self.assertEqual(document.tables[1].cell(1, 1).text, "Klaim contoh")
        self.assertEqual(document.tables[2].cell(1, 4).text, "sesuai")
        self.assertIn("Perlu bukti tambahan.", [p.text for p in document.paragraphs])

    def test_personnel_zero_values_and_none_years(self):
        target = self.root / "nested" / "personnel.docx"
        payload = {"personel": [{"nik": "000000", "pengalaman_kerja_bulan": 0,
                                 "pengalaman_kerja_tahun": None, "pengalaman_min_kak_tahun": 0}]}
        cv_report.render_personnel_docx(payload, target)
        cells = Document(target).tables[0].rows[1].cells
        self.assertEqual([cell.text for cell in cells], ["1", "", "000000", "", "", "", "0 Tahun", "0", ""])

    def test_empty_review_and_scalar_biodata_remain_supported(self):
        payload = {"cv_file": "test.pdf", "biodata": "Belum dipetakan",
                   "experience_validation": None, "attachment_cross_check": [], "findings": []}
        target = self.root / "empty.docx"
        result = cv_report.render_report(payload, target)
        document = Document(target)
        self.assertEqual(result["finding_count"], 0)
        self.assertEqual(document.tables[0].cell(1, 1).text, "Belum dipetakan")
        self.assertIn("Tidak ada temuan.", [p.text for p in document.paragraphs])

    def test_invalid_review_is_rejected_before_writing(self):
        target = self.root / "invalid.docx"
        for payload in ([], {}, {**REVIEW, "findings": "not a list"},
                        {**REVIEW, "experience_validation": ["not an object"]},
                        {**REVIEW, "internet_sources": {}}):
            with self.subTest(payload=payload), self.assertRaises(ValueError):
                cv_report.render_report(payload, target)
            self.assertFalse(target.exists())

    def test_cli_from_clean_directory_and_utf8_stdin(self):
        # No templates directory, models, or project-specific scripts are copied.
        for script_name, payload, suffix, marker, extra in (
            ("exporter.py", PERSONNEL, "xlsx", "HERMES_EXPORT_RESULT=", ["--template", "daftar_tenaga_ahli"]),
            ("cv_report.py", REVIEW, "docx", "HERMES_CV_REPORT_RESULT=", []),
        ):
            with self.subTest(script=script_name):
                script = self.root / script_name
                shutil.copyfile(Path(__file__).parent / script_name, script)
                data = copy.deepcopy(payload)
                if suffix == "xlsx":
                    data["personel"][0]["nama_personel"] = "Contoh 日本 & é"
                else:
                    data["biodata"]["nama"] = "Contoh 日本 & é"
                target = self.root / f"stdin.{suffix}"
                env = {**os.environ, "PYTHONIOENCODING": "cp1252"}
                completed = subprocess.run(
                    [sys.executable, str(script), *extra, "--payload-json", "-", "--output", str(target)],
                    input=json.dumps(data, ensure_ascii=False), text=True, encoding="utf-8",
                    capture_output=True, cwd=self.root, env=env, timeout=30,
                )
                self.assertEqual(completed.returncode, 0, completed.stderr + completed.stdout)
                result = json.loads(completed.stdout.split(marker)[1])
                self.assertTrue(result["success"])
                if suffix == "xlsx":
                    workbook = load_workbook(target)
                    self.assertEqual(workbook.active["B6"].value, "Contoh 日本 & é")
                    workbook.close()
                else:
                    self.assertEqual(Document(target).tables[0].cell(1, 1).text, "Contoh 日本 & é")
                target.unlink()
                failed = subprocess.run(
                    [sys.executable, str(script), *extra, "--payload-json", "{", "--output", str(target)],
                    text=True, encoding="utf-8", capture_output=True, cwd=self.root, timeout=30,
                )
                self.assertEqual(failed.returncode, 1)
                self.assertFalse(json.loads(failed.stdout.split(marker)[1])["success"])
                self.assertFalse(target.exists())
        self.assertFalse((self.root / "templates").exists())


if __name__ == "__main__":
    unittest.main()
