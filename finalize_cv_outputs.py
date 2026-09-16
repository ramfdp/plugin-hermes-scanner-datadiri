import json, re
from pathlib import Path
from copy import deepcopy


from openpyxl import load_workbook
from exporter import export_payload
from cv_report import render_personnel_docx

OUT_DIR = Path(r"C:\hscan\cv_outputs")
payload_path = OUT_DIR / "payload_daftar_tenaga_ahli.json"
payload = json.loads(payload_path.read_text(encoding="utf-8"))

# Fix fields found during QA.
for p in payload["personel"]:
    if p["nama_personel"] == "Anderson Hario Pangestiaji" and p["kualifikasi_pendidikan"] == "2011":
        p["kualifikasi_pendidikan"] = "S1 Teknik Sipil Universitas Atma Jaya Yogyakarta 2011"
    if p["nama_personel"] == "Dudung Jatnika":
        p["kualifikasi_pendidikan"] = p["kualifikasi_pendidikan"].replace("i S1", "S1").strip()
    if p["nama_personel"] == "Tubagus Pirmannudin":
        p["kualifikasi_pendidikan"] = p["kualifikasi_pendidikan"].replace('*"', '').strip()
    if p["nama_personel"] == "Ahmad Idris Sirojudin":
        p["kualifikasi_pendidikan"] = p["kualifikasi_pendidikan"].replace('*.', '').strip()
    if p["nama_personel"] == "Agung Satria Putra" and p["pengalaman_kerja_tahun"]:
        p["pengalaman_kerja_tahun"] = round(float(p["pengalaman_kerja_bulan"])/12, 2)
    # Two KTP images were too small for line OCR; values read from high-res page images.
    if p["nama_personel"] == "Rizal Muttaqin" and not p["nik"]:
        p["nik"] = "3514130104860001"
    if p["nama_personel"] == "Ari Artanto" and not p["nik"]:
        p["nik"] = "3305150408750002"

payload_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

xlsx_path = OUT_DIR / "Daftar_Tenaga_Ahli_Jawa_1.xlsx"
docx_path = OUT_DIR / "Daftar_Tenaga_Ahli_Jawa_1.docx"

# Excel via cloned exporter script implementation.
result = export_payload("daftar_tenaga_ahli", payload, xlsx_path)

# Word table output using python-docx dependency shipped by the same scanner repo.
render_personnel_docx(payload, docx_path)

# Verification readback
wb = load_workbook(xlsx_path, read_only=True, data_only=True)
ws = wb.active
row_count = max(0, ws.max_row - 5)
nonempty_names = sum(1 for r in range(6, ws.max_row+1) if ws.cell(r,2).value)
wb.close()

print(json.dumps({
    'success': True,
    'xlsx': str(xlsx_path),
    'docx': str(docx_path),
    'payload': str(payload_path),
    'excel_rows': row_count,
    'nonempty_names': nonempty_names,
    'payload_rows': len(payload['personel']),
    'missing_nik': sum(1 for p in payload['personel'] if not p.get('nik')),
    'exporter_result': result,
}, ensure_ascii=False, indent=2))
