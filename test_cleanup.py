import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.enum.section import WD_ORIENT
from cv_report import render_personnel_docx
from tools import _child_env


class CleanupTest(unittest.TestCase):
    def test_personnel_word_output(self):
        with tempfile.TemporaryDirectory() as temp:
            target = Path(temp) / "personnel.docx"
            render_personnel_docx({"personel": [{}, {
                "nama_personel": "Example & Test", "nik": "001234",
                "pengalaman_kerja_tahun": 1.25, "pengalaman_min_kak_tahun": 0,
                "pengalaman_kerja_bulan": 15,
            }]}, target)
            document = Document(target)
            self.assertEqual(document.sections[0].orientation, WD_ORIENT.LANDSCAPE)
            self.assertEqual(document.paragraphs[0].text, "DAFTAR TENAGA AHLI")
            rows = document.tables[0].rows
            self.assertEqual([cell.text for cell in rows[1].cells], ["1"] + [""] * 8)
            self.assertEqual([cell.text for cell in rows[2].cells],
                             ["2", "Example & Test", "001234", "", "", "", "0 Tahun", "15", "1,25"])

    def test_child_environment_is_isolated(self):
        original = {"PATH": "original", "PYTHONPATH": "old", "PYTHONHOME": "old",
                    "VIRTUAL_ENV": "old", "KEEP": "value"}
        with patch.dict(os.environ, original, clear=True):
            env = _child_env(Path("runtime"))
            self.assertEqual(env, {"PATH": str(Path("runtime/.venv/Scripts")) + os.pathsep + "original",
                                   "KEEP": "value"})
            self.assertEqual(dict(os.environ), original)


if __name__ == "__main__":
    unittest.main()
